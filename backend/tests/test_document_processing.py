import os
import asyncio
import unittest
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("DATABASE_URL", "sqlite://")

from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from fastapi import HTTPException
from starlette.background import BackgroundTasks
from starlette.requests import Request
from starlette.datastructures import UploadFile
from openpyxl import load_workbook
from PIL import Image

from backend.app import main
from backend.app.database import Base
from backend.app.models import Document, Lead, LeadContact, LeadIdempotency, LeadSourceImage, User
from backend.app.schemas import LeadCreate


class DocumentReviewFlowTests(unittest.TestCase):
    def test_lead_idempotency_is_persistent_and_scoped_to_user(self):
        with self.session_factory() as db:
            first_user = User(full_name="First", email="first@example.com", normalized_email="first@example.com", password_hash="x")
            second_user = User(full_name="Second", email="second@example.com", normalized_email="second@example.com", password_hash="x")
            db.add_all([first_user, second_user])
            db.commit()
            db.refresh(first_user)
            db.refresh(second_user)
            request = Request({"type": "http", "method": "POST", "path": "/api/leads", "headers": [(b"x-idempotency-key", b"same-key")]})
            first = main.create_lead(LeadCreate(company_name="Acme", contact_numbers=["9876543210"]), request, db, first_user)
            replay = main.create_lead(LeadCreate(company_name="Acme", contact_numbers=["9876543210"]), request, db, first_user)
            other = main.create_lead(LeadCreate(company_name="Acme", contact_numbers=["9876543210"]), request, db, second_user)
            self.assertEqual(first.id, replay.id)
            self.assertNotEqual(first.id, other.id)
            self.assertEqual(db.scalar(select(func.count()).select_from(LeadIdempotency)), 2)

    def test_text_preview_preserves_original_and_multiple_contacts_without_company(self):
        source = 'Contacts\r\n+91 89995 21798\r\n+91 74994 22904\r\nपुणे\r\n'
        with tempfile.SpooledTemporaryFile() as stream:
            stream.write(source.encode('utf-8-sig'))
            stream.seek(0)
            upload = UploadFile(filename='contacts.txt', file=stream)
            # This is an import-flow unit test, not a multi-gigabyte local
            # translation-model integration test.
            with patch("backend.app.ocr.translate_document_text", return_value=(None, "unavailable")), \
                 patch("backend.app.translation.translate_document_text", return_value=(None, "unavailable")):
                result = asyncio.run(main.preview_raw_leads(upload))
        fields = result['candidates'][0]['fields']
        self.assertEqual(fields['original_ocr_text'], source)
        self.assertEqual(fields['raw_text'], source)
        self.assertEqual(fields['contact_numbers'], ['+91 89995 21798', '+91 74994 22904'])

    def setUp(self) -> None:
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(bind=self.engine)
        self.session_factory = sessionmaker(bind=self.engine, autoflush=False, autocommit=False)
        with self.session_factory() as db:
            db.add(Document(
                id="review-document",
                original_filename="business-card.png",
                content_type="image/png",
                storage_path="unused-for-mocked-ocr.png",
                processing_status="uploaded",
            ))
            db.commit()

    def tearDown(self) -> None:
        self.engine.dispose()

    def test_untranslated_ocr_cannot_be_saved_as_english(self) -> None:
        with self.session_factory() as db:
            with self.assertRaises(HTTPException) as error:
                main.create_lead(LeadCreate(company_name="अनमोल", contact_numbers=["9876543210"], document_id="review-document", source_review_confirmed=True), db)
            self.assertEqual(error.exception.status_code, 422)
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 0)

    def test_english_save_keeps_original_transcript_and_export_values(self) -> None:
        from backend.app.schemas import LeadUpdate
        with self.session_factory() as db:
            doc = db.get(Document, "review-document")
            doc.extracted_fields = {"original_ocr_text": "अनमोल", "raw_text": "Anmol"}
            db.commit()
            saved = main.create_lead(LeadCreate(company_name="Anmol", contact_numbers=["9876543210"], raw_text="Anmol", document_id=doc.id, source_review_confirmed=True), db)
            self.assertEqual(saved.original_ocr_text, "अनमोल")
            self.assertEqual(saved.company_name, "Anmol")
            with self.assertRaises(HTTPException):
                main.update_lead(saved.id, LeadUpdate(company_name="अनमोल", source_review_confirmed=True), db)
            self.assertEqual(db.get(Lead, saved.id).company_name, "Anmol")
            class CapturedResponse:
                def __init__(self, content, **kwargs):
                    self.content = content.getvalue()
            with patch.object(main, "StreamingResponse", CapturedResponse):
                exported = main.export_leads(search=None, ids=[saved.id], company_name=None,
                    owner_name=None, city=None, state=None, export_count=None, db=db)
            sheet = load_workbook(BytesIO(exported.content))["Leads"]
            self.assertEqual(sheet["C2"].value, "Anmol")

    def test_business_description_is_saved_updated_and_exported_separately(self) -> None:
        from backend.app.schemas import LeadUpdate
        with self.session_factory() as db:
            lead = main.create_lead(LeadCreate(
                company_name="SM Electricals",
                contact_numbers=["9876543210"],
                business_description="Lighting installation and maintenance",
                raw_text="Company Name: SM Electricals\nDescription: Lighting installation and maintenance",
                source_review_confirmed=True,
            ), db)
            self.assertEqual(lead.business_description, "Lighting installation and maintenance")
            updated = main.update_lead(lead.id, LeadUpdate(
                business_description="Electrical repair and installation", source_review_confirmed=True,
            ), db)
            self.assertEqual(updated.business_description, "Electrical repair and installation")

            class CapturedResponse:
                def __init__(self, content, **kwargs): self.content = content.getvalue()
            with patch.object(main, "StreamingResponse", CapturedResponse):
                exported = main.export_leads(search=None, ids=[lead.id], company_name=None,
                    owner_name=None, city=None, state=None, export_count=None, db=db)
            sheet = load_workbook(BytesIO(exported.content))["Leads"]
            description_column = [cell.value for cell in sheet[1]].index("Business description") + 1
            self.assertEqual(sheet.cell(2, description_column).value, "Electrical repair and installation")

    def test_original_script_in_other_document_text_does_not_block_save(self) -> None:
        with self.session_factory() as db:
            saved = main.create_lead(
                LeadCreate(
                    company_name="Anmol Boutique",
                    contact_numbers=["9876543210"],
                    owner_name="Anmol Verma",
                    designation="Owner",
                    raw_text="अनमोल बुटीक\nसभी प्रकार के कपड़े",
                    original_ocr_text="अनमोल बुटीक\nसभी प्रकार के कपड़े",
                    document_id="review-document",
                    source_review_confirmed=True,
                ),
                db,
            )
            self.assertEqual(saved.raw_text, "अनमोल बुटीक\nसभी प्रकार के कपड़े")
            self.assertEqual(saved.original_ocr_text, "अनमोल बुटीक\nसभी प्रकार के कपड़े")

    def test_ocr_keeps_extracted_fields_as_review_draft(self) -> None:
        extracted_fields = {
            "raw_text": "Acme Solutions Pvt Ltd\nJane Doe\nDirector\n211 Main Road",
            "cards": [{
                "company_name": "Acme Solutions Pvt Ltd",
                "owner_name": "Jane Doe",
                "designation": "Director",
                "company_address": "211 Main Road",
                "contact_numbers": ["+1 212 555 0100", "+1 212 555 0101"],
                "email_addresses": ["info@acme.example", "sales@acme.example"],
                "raw_text": "Acme Solutions Pvt Ltd\nJane Doe\nDirector\n211 Main Road",
            }],
        }

        with patch.object(main, "SessionLocal", self.session_factory), patch.object(
            main, "extract_fields", return_value=extracted_fields
        ):
            main._process_document("review-document")

        with self.session_factory() as db:
            document = db.get(Document, "review-document")
            self.assertEqual(document.processing_status, "needs_review")
            self.assertIsNone(document.error_message)
            self.assertEqual(document.extracted_fields, extracted_fields)
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 0)

    def test_missing_company_name_stays_as_manual_review_draft(self) -> None:
        extracted_fields = {"raw_text": "+1 212 555 0100", "cards": [{"contact_numbers": ["+1 212 555 0100"]}]}
        with patch.object(main, "SessionLocal", self.session_factory), patch.object(
            main, "extract_fields", return_value=extracted_fields
        ):
            main._process_document("review-document")

        with self.session_factory() as db:
            document = db.get(Document, "review-document")
            self.assertEqual(document.processing_status, "needs_review")
            self.assertIn("enter the details manually", document.error_message)
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 0)

    def test_ocr_failure_keeps_document_retryable_without_creating_a_lead(self) -> None:
        with patch.object(main, "SessionLocal", self.session_factory), patch.object(
            main, "extract_fields", side_effect=RuntimeError("OCR failed")
        ):
            main._process_document("review-document")

        with self.session_factory() as db:
            document = db.get(Document, "review-document")
            self.assertEqual(document.processing_status, "failed")
            self.assertTrue(document.error_message)
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 0)

    def test_retry_only_claims_failed_scan_once_and_keeps_source(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
            handle.write(b"source")
            source_path = Path(handle.name)
        self.addCleanup(source_path.unlink, missing_ok=True)
        with self.session_factory() as db:
            document = db.get(Document, "review-document")
            document.processing_status = "failed"
            document.storage_path = str(source_path)
            db.commit()
            request = Request({"type": "http", "method": "POST", "path": "/api/documents/review-document/retry", "headers": []})
            request.state.request_id = "request-id"
            tasks = BackgroundTasks()
            import threading
            slots = threading.BoundedSemaphore(1)
            with patch.object(main, "_ocr_slots", slots):
                queued = main.retry_document("review-document", request, tasks, db)
                self.assertEqual(queued.processing_status, "uploaded")
                self.assertTrue(source_path.is_file())
                with self.assertRaises(HTTPException) as error:
                    main.retry_document("review-document", request, BackgroundTasks(), db)
                self.assertEqual(error.exception.status_code, 409)
                self.assertEqual(len(tasks.tasks), 1)
                self.assertFalse(slots.acquire(blocking=False))
                slots.release()

    def test_user_confirmation_is_required_before_ocr_draft_is_saved(self) -> None:
        with self.session_factory() as db:
            draft = LeadCreate(company_name="Acme Solutions Pvt Ltd", contact_numbers=["1111111111"], document_id="review-document")
            with self.assertRaises(HTTPException) as error:
                main.create_lead(draft, db)
            self.assertEqual(error.exception.status_code, 422)
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 0)

            confirmed = LeadCreate(
                company_name="Acme Solutions Pvt Ltd",
                document_id="review-document",
                contact_numbers=["1111111111", "2222222222"],
                email_addresses=["one@acme.example", "two@acme.example"],
                source_review_confirmed=True,
            )
            saved = main.create_lead(confirmed, db)
            self.assertEqual(saved.verification_status, "verified")
            self.assertEqual(saved.contact_numbers, ["+91 1111111111", "+91 2222222222"])
            self.assertEqual(saved.email_addresses, ["one@acme.example", "two@acme.example"])

    def test_indian_mobile_numbers_are_normalized_when_saved(self) -> None:
        with self.session_factory() as db:
            saved = main.create_lead(LeadCreate(
                company_name="Acme Solutions",
                contact_numbers=["9876543210", "022 41310098", "+1 212 555 0100"],
            ), db)
            self.assertEqual(saved.contact_numbers, ["+91 9876543210", "022 41310098", "+1 212 555 0100"])

    def test_editing_lead_persists_multiple_contacts_and_emails(self) -> None:
        with self.session_factory() as db:
            saved = main.create_lead(LeadCreate(company_name="Acme Pvt Ltd", contact_numbers=["+91 98765 00000"]), db)
            updated = main.update_lead(saved.id, main.LeadUpdate(
                contact_numbers=["+91 98765 43210", "+91 98765 43211"],
                email_addresses=["info@acme.example", "sales@acme.example"],
            ), db)
            self.assertEqual(updated.contact_numbers, ["+91 98765 43210", "+91 98765 43211"])
            self.assertEqual(updated.email_addresses, ["info@acme.example", "sales@acme.example"])

    def test_leads_list_includes_separate_contact_values(self) -> None:
        with self.session_factory() as db:
            saved = main.create_lead(LeadCreate(
                company_name="Acme Pvt Ltd",
                contact_numbers=["+91 98765 43210", "+91 98765 43211"],
                email_addresses=["info@acme.example", "sales@acme.example"],
            ), db)
            response = main.list_leads(page=1, page_size=50, db=db)
            listed = response["leads"][0]

            self.assertEqual(listed.id, saved.id)
            self.assertEqual(listed.contact_numbers, ["+91 98765 43210", "+91 98765 43211"])
            self.assertEqual(listed.email_addresses, ["info@acme.example", "sales@acme.example"])

    def test_front_back_sources_are_linked_to_one_explicitly_saved_lead(self) -> None:
        with self.session_factory() as db:
            db.add(Document(id="back-document", original_filename="back.png", content_type="image/png", storage_path="back.png", processing_status="needs_review", extracted_fields={"raw_text": "sales@acme.example"}))
            db.commit()
            saved = main.create_lead(LeadCreate(
                company_name="Acme Solutions Pvt Ltd",
                contact_numbers=["9876543210"],
                document_id="review-document",
                source_document_ids=["review-document", "back-document"],
                source_sides=["front", "back"],
                facebook_link="https://facebook.com/acme",
                instagram_link="https://instagram.com/acme",
                whatsapp_link="https://wa.me/919876543210",
                source_review_confirmed=True,
            ), db)
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 1)
            self.assertEqual(db.scalar(select(func.count()).select_from(LeadSourceImage)), 2)
            self.assertEqual(saved.source_document_ids, ["review-document", "back-document"])
            self.assertEqual([image.side for image in saved.source_images], ["front", "back"])
            listed = main.list_leads(page=1, page_size=50, db=db)["leads"][0]
            self.assertEqual(listed.facebook_link, "https://facebook.com/acme")
            self.assertEqual(listed.instagram_link, "https://instagram.com/acme")
            self.assertEqual(listed.whatsapp_link, "https://wa.me/919876543210")
            self.assertEqual(listed.source_document_ids, ["review-document", "back-document"])

    def test_batch_generate_creates_all_rows_and_keeps_them_marked_for_review(self) -> None:
        payload = [
            main.LeadFields(
                company_name="Batch One Pvt Ltd",
                contact_numbers=["9876543210", "9123456789"],
                email_addresses=["one@example.test", "sales@example.test"],
                raw_text="Imported spreadsheet row",
            ),
            main.LeadFields(company_name="Batch Two Ltd", contact_numbers=["9988776655"]),
        ]
        with self.session_factory() as db:
            result = main.create_leads_batch(payload, db)
            self.assertEqual(result["created"], 2)
            saved = db.scalar(select(Lead).where(Lead.company_name == "Batch One Pvt Ltd"))
            self.assertEqual(saved.verification_status, "needs_review")
            self.assertEqual(saved.contact_numbers, ["+91 9876543210", "+91 9123456789"])
            self.assertEqual(saved.email_addresses, ["one@example.test", "sales@example.test"])

    def test_legacy_ten_digit_phone_is_prefixed_in_lead_list_and_excel_export(self) -> None:
        class CapturedResponse:
            def __init__(self, content, **_kwargs):
                self.content = content.getvalue()

        with self.session_factory() as db:
            saved = main.create_lead(LeadCreate(company_name="Legacy Contact", contact_numbers=["9876543210"]), db)
            saved.contacts[0].value = "9876543210"  # Simulate an existing pre-normalization database row.
            db.commit()
            listed = main.list_leads(page=1, page_size=50, db=db)["leads"][0]
            self.assertEqual(listed.contact_numbers, ["+91 9876543210"])
            with patch.object(main, "StreamingResponse", CapturedResponse):
                exported = main.export_leads(
                    search=None, ids=[], company_name=None, owner_name=None, city=None, state=None,
                    export_count=None, db=db,
                )
        workbook = load_workbook(BytesIO(exported.content))
        self.assertEqual(workbook["Leads"]["H2"].value, "+91 9876543210")

    def test_raw_data_headers_map_values_to_separate_fields(self) -> None:
        candidate = main._raw_row_fields(
            ["Company Name", "Owner Name", "Designation", "Company Address", "Mobile 1", "Mobile 2", "Email 1", "Email 2", "Website"],
            ["Acme Pvt Ltd", "Jane Doe", "Director", "4 Main Road", "9876543210", "9123456789", "jane@acme.test", "sales@acme.test", "acme.test"],
        )
        self.assertEqual(candidate["company_name"], "Acme Pvt Ltd")
        self.assertEqual(candidate["contact_numbers"], ["9876543210", "9123456789"])
        self.assertEqual(candidate["email_addresses"], ["jane@acme.test", "sales@acme.test"])

    def test_raw_excel_common_aliases_and_numeric_phone_cells_are_mapped_safely(self) -> None:
        candidate = main._raw_row_fields(
            ["Business", "Primary Contact Name", "Job Title", "Office Location", "Mobile Number 1", "Mobile Number 2", "Email Address 1", "Email Address 2", "Website URL", "Revenue", "Email Verification Status", "Emergency Contact Number"],
            ["Acme Foods Pvt Ltd", "Jane Doe", "Director", "12 Sample Road", 9876543210.0, 9123456789.0, "jane@acme.example", "sales@acme.example", "acme.example", "₹50 lakh", "Verified", "911234567890"],
        )

        self.assertEqual(candidate["company_name"], "Acme Foods Pvt Ltd")
        self.assertEqual(candidate["owner_name"], "Jane Doe")
        self.assertEqual(candidate["designation"], "Director")
        self.assertEqual(candidate["company_address"], "12 Sample Road")
        self.assertEqual(candidate["contact_numbers"], ["9876543210", "9123456789"])
        self.assertEqual(candidate["email_addresses"], ["jane@acme.example", "sales@acme.example"])
        self.assertEqual(candidate["website"], "acme.example")
        self.assertNotIn("Revenue", candidate["raw_text"])
        self.assertNotIn("Email Verification Status", candidate["raw_text"])
        self.assertNotIn("Emergency Contact Number", candidate["raw_text"])

    def test_raw_excel_ignores_columns_not_present_on_add_lead_form(self) -> None:
        candidate = main._raw_row_fields(
            ["Company Name", "Owner Name", "Contact Number 1", "Email ID 1", "Other Document Text", "Revenue", "Internal Note"],
            ["Acme Pvt Ltd", "Jane Doe", "9876543210", "jane@acme.test", "Trade fair contact", "₹50 lakh", "Do not import"],
        )
        self.assertEqual(candidate["company_name"], "Acme Pvt Ltd")
        self.assertEqual(candidate["owner_name"], "Jane Doe")
        self.assertEqual(candidate["contact_numbers"], ["9876543210"])
        self.assertEqual(candidate["email_addresses"], ["jane@acme.test"])
        self.assertIn("Other Document Text: Trade fair contact", candidate["raw_text"])
        self.assertNotIn("Revenue", candidate["raw_text"])
        self.assertNotIn("Internal Note", candidate["raw_text"])

    def test_raw_excel_finds_headers_after_blank_rows_and_maps_company_work_data(self) -> None:
        rows = [
            [None, None, None, None, None, None],
            [None, None, None, None, None, None],
            [None, "Company Name", "Contact No", "Email Id", "Website", "Company Work Data"],
            [1, "Acme Logistics Pvt Ltd", 9876543210.0, "info@acme.example", "acme.example", "Provides freight and warehousing services."],
        ]

        candidates = main._raw_candidates_from_rows(rows)

        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0]["row"], 4)
        fields = candidates[0]["fields"]
        self.assertEqual(fields["company_name"], "Acme Logistics Pvt Ltd")
        self.assertEqual(fields["contact_numbers"], ["9876543210"])
        self.assertEqual(fields["email_addresses"], ["info@acme.example"])
        self.assertEqual(fields["website"], "acme.example")
        self.assertIn("Company Work Data: Provides freight and warehousing services.", fields["raw_text"])

    def test_sample_excel_import_saves_ten_leads_with_separate_contacts(self) -> None:
        sample_path = Path(__file__).resolve().parents[2] / "test-data" / "sample-leads-10.xlsx"
        workbook = load_workbook(sample_path, read_only=True, data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
        headers = [str(value or "").strip() for value in rows[0]]
        candidates = [
            main._raw_row_fields(headers, list(row))
            for row in rows[1:]
            if any(value not in (None, "") for value in row)
        ]

        with self.session_factory() as db:
            saved = [main.create_lead(LeadCreate(
                company_name=candidate["company_name"],
                owner_name=candidate["owner_name"],
                designation=candidate["designation"],
                company_address=candidate["company_address"],
                contact_numbers=candidate["contact_numbers"],
                email_addresses=candidate["email_addresses"],
                website=candidate["website"],
                raw_text=candidate["raw_text"],
                source_review_confirmed=True,
            ), db) for candidate in candidates]
            self.assertEqual(len(saved), 10)
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), 10)
            self.assertEqual(db.scalar(select(func.count()).select_from(LeadContact)), 30)
            self.assertEqual(sum(len(lead.contact_numbers) == 2 for lead in saved), 5)
            self.assertEqual(sum(len(lead.email_addresses) == 2 for lead in saved), 5)

    def test_excel_export_sizes_columns_and_wraps_long_cells(self) -> None:
        company_name = "Acme Long Company Name Pvt Ltd with a Detailed Legal Suffix"
        address = "Unit 1, Sample Industrial Area, Sector 25, Navi Mumbai, Maharashtra, India 400705"
        with self.session_factory() as db:
            main.create_lead(LeadCreate(
                company_name=company_name,
                company_address=address,
                contact_numbers=["9876543210", "9123456789"],
                email_addresses=["contact@acme.example", "sales@acme.example"],
                facebook_link="https://facebook.com/acme",
                instagram_link="https://instagram.com/acme",
                whatsapp_link="https://wa.me/919876543210",
                raw_text="Company Name: " + company_name + "\nAddress: " + address,
                source_review_confirmed=True,
            ), db)
            class CapturedResponse:
                def __init__(self, content, **_kwargs):
                    self.content = content.getvalue()

            with patch.object(main, "StreamingResponse", CapturedResponse):
                response = main.export_leads(
                    search=None, ids=[], company_name=None, owner_name=None, city=None, state=None,
                    export_count=None, db=db,
                )
            workbook_bytes = response.content

        workbook = load_workbook(BytesIO(workbook_bytes))
        sheet = workbook["Leads"]
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["A1"].font.bold, True)
        self.assertEqual(sheet["A1"].value, "S.No.")
        self.assertEqual(sheet["A2"].value, 1)
        self.assertEqual(sheet["B1"].value, "Autogenerated ID")
        self.assertEqual(sheet["H1"].value, "Contact Numbers")
        self.assertEqual(sheet["I1"].value, "Email ID 1")
        self.assertEqual(sheet["H2"].value, "+91 9876543210\n+91 9123456789")
        self.assertEqual(sheet["I2"].value, "contact@acme.example")
        self.assertEqual(sheet["J2"].value, "sales@acme.example")
        self.assertEqual(sheet["L1"].value, "Facebook link")
        self.assertEqual(sheet["L2"].value, "https://facebook.com/acme")
        self.assertEqual(sheet["M2"].value, "https://instagram.com/acme")
        self.assertEqual(sheet["N2"].value, "https://wa.me/919876543210")
        self.assertGreaterEqual(sheet.column_dimensions["C"].width, min(len(company_name) + 2, 60))
        self.assertEqual(sheet["C2"].value, company_name)
        self.assertTrue(sheet["E2"].alignment.wrap_text)
        self.assertGreater(sheet.row_dimensions[2].height, 20)

    def test_lead_filters_intersect_city_and_state_with_existing_address(self) -> None:
        with self.session_factory() as db:
            main.create_lead(LeadCreate(
                company_name="Acme Pvt Ltd", owner_name="Jane Doe", company_address="Navi Mumbai, Maharashtra",
                contact_numbers=["9876543201"],
            ), db)
            main.create_lead(LeadCreate(
                company_name="Acme Pvt Ltd", owner_name="John Doe", company_address="Pune, Maharashtra",
                contact_numbers=["9876543202"],
            ), db)
            statement = main._filter_lead_statement(
                select(Lead), company_name="Acme", owner_name="Jane", city="Mumbai", state="Maharashtra",
            )
            matched = db.scalars(statement).all()
            self.assertEqual(len(matched), 1)
            self.assertEqual(matched[0].owner_name, "Jane Doe")

    def test_zip_creates_one_pending_document_per_supported_card_file(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            archive_bytes = BytesIO()
            png = BytesIO()
            Image.new("RGB", (2, 2), "white").save(png, format="PNG")
            jpeg = BytesIO()
            Image.new("RGB", (2, 2), "white").save(jpeg, format="JPEG")
            with zipfile.ZipFile(archive_bytes, "w") as archive:
                archive.writestr("cards/first.png", png.getvalue())
                archive.writestr("cards/second.jpg", jpeg.getvalue())
                archive.writestr("notes.txt", "skip")
            archive_bytes.seek(0)
            with self.session_factory() as db, patch.object(main, "UPLOAD_DIR", Path(temporary_directory)):
                documents = main._documents_from_zip(db, "cards.zip", archive_bytes)
                self.assertEqual([document.original_filename for document in documents], ["first.png", "second.jpg"])
                self.assertTrue(all(Path(document.storage_path).is_file() for document in documents))
                db.rollback()

    def test_duplicate_phone_or_email_blocks_save(self) -> None:
        with self.session_factory() as db:
            first = main.create_lead(LeadCreate(
                company_name="First Corp",
                contact_numbers=["+91 9876543210"],
                email_addresses=["info@first.example"],
            ), db)
            self.assertIsNotNone(first.id)

            # Same phone (unprefixed) should be blocked with 409
            with self.assertRaises(HTTPException) as ctx:
                main.create_lead(LeadCreate(
                    company_name="Second Corp",
                    contact_numbers=["9876543210"],
                ), db)
            self.assertEqual(ctx.exception.status_code, 409)
            self.assertIn("Duplicate lead detected", ctx.exception.detail)

            # Same email (different case) should be blocked with 409
            with self.assertRaises(HTTPException) as ctx:
                main.create_lead(LeadCreate(
                    company_name="Third Corp",
                    contact_numbers=["9333333333"],
                    email_addresses=["INFO@FIRST.EXAMPLE"],
                ), db)
            self.assertEqual(ctx.exception.status_code, 409)
            self.assertIn("Duplicate lead detected", ctx.exception.detail)

            # Check duplicate endpoint
            check_req = main.DuplicateCheckRequest(contact_numbers=["9876543210"])
            check_res = main.check_duplicate_lead(check_req, db, user=None)
            self.assertTrue(check_res.is_duplicate)
            self.assertEqual(check_res.existing_company_name, "First Corp")

            # Check duplicate endpoint with exclude_lead_id
            check_res_excluded = main.check_duplicate_lead(
                main.DuplicateCheckRequest(contact_numbers=["9876543210"], exclude_lead_id=first.id),
                db, user=None,
            )
            self.assertFalse(check_res_excluded.is_duplicate)

    def test_batch_import_skips_duplicates(self) -> None:
        with self.session_factory() as db:
            main.create_lead(LeadCreate(
                company_name="Existing Corp",
                contact_numbers=["9876543210"],
            ), db)

            payload = [
                main.LeadFields(company_name="Row 1", contact_numbers=["9876543210"]),
                main.LeadFields(company_name="Row 2", contact_numbers=["9111111111"]),
                main.LeadFields(company_name="Row 3", contact_numbers=["9111111111"]),
                main.LeadFields(company_name="Row 4", contact_numbers=["9222222222"]),
            ]
            res = main.create_leads_batch(payload, db)
            self.assertEqual(res["created"], 2)
            self.assertEqual(res["skipped_duplicates"], 2)

    def test_export_leads_custom_file_name_and_sanitization(self) -> None:
        with self.session_factory() as db:
            lead = main.create_lead(LeadCreate(
                company_name="Test Company",
                contact_numbers=["9876543210"],
            ), db)

            class CapturedResponse:
                def __init__(self, content, headers=None, **_kwargs):
                    self.content = content.getvalue()
                    self.headers = headers or {}

            # 1. Default filename when omitted
            with patch.object(main, "StreamingResponse", CapturedResponse):
                res1 = main.export_leads(ids=[lead.id], db=db)
            self.assertEqual(res1.headers.get("Content-Disposition"), 'attachment; filename="leads.xlsx"')

            # 2. Custom human-readable filename
            with patch.object(main, "StreamingResponse", CapturedResponse):
                res2 = main.export_leads(ids=[lead.id], filename="October leads", db=db)
            self.assertEqual(res2.headers.get("Content-Disposition"), 'attachment; filename="October leads.xlsx"')

            # 3. Path traversal attempt sanitized
            with patch.object(main, "StreamingResponse", CapturedResponse):
                res3 = main.export_leads(ids=[lead.id], filename="../../secret-report", db=db)
            self.assertEqual(res3.headers.get("Content-Disposition"), 'attachment; filename="secret-report.xlsx"')

            # 4. Invalid characters sanitized
            with patch.object(main, "StreamingResponse", CapturedResponse):
                res4 = main.export_leads(ids=[lead.id], filename='leads:2026*final?"test<file>|', db=db)
            self.assertEqual(res4.headers.get("Content-Disposition"), 'attachment; filename="leads2026finaltestfile.xlsx"')

            # 5. Duplicate .xlsx extension avoided
            with patch.object(main, "StreamingResponse", CapturedResponse):
                res5 = main.export_leads(ids=[lead.id], filename="my-report.xlsx.xlsx", db=db)
            self.assertEqual(res5.headers.get("Content-Disposition"), 'attachment; filename="my-report.xlsx"')

            # 6. Very long name capped at 120 characters
            long_name = "X" * 200
            with patch.object(main, "StreamingResponse", CapturedResponse):
                res6 = main.export_leads(ids=[lead.id], filename=long_name, db=db)
            expected_header = f'attachment; filename="{"X" * 120}.xlsx"'
            self.assertEqual(res6.headers.get("Content-Disposition"), expected_header)

            # 7. Exactly one 'Contact Numbers' column in exported workbook
            workbook = load_workbook(BytesIO(res2.content))
            sheet = workbook["Leads"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(headers.count("Contact Numbers"), 1)
            self.assertEqual(headers.count("Contact Number 1"), 0)

    def test_export_leads_selected_rows_and_filters(self) -> None:
        with self.session_factory() as db:
            l1 = main.create_lead(LeadCreate(company_name="Alpha", owner_name="Alice", company_address="Pune, Maharashtra", contact_numbers=["9876543201"]), db)
            l2 = main.create_lead(LeadCreate(company_name="Beta", owner_name="Bob", company_address="Mumbai, Maharashtra", contact_numbers=["9876543202"]), db)
            l3 = main.create_lead(LeadCreate(company_name="Gamma", owner_name="Charlie", company_address="Delhi", contact_numbers=["9876543203"]), db)

            class CapturedResponse:
                def __init__(self, content, headers=None, **_kwargs):
                    self.content = content.getvalue()
                    self.headers = headers or {}

            # Selected IDs take precedence
            with patch.object(main, "StreamingResponse", CapturedResponse):
                res = main.export_leads(ids=[l1.id, l3.id], city="Mumbai", db=db)
            wb = load_workbook(BytesIO(res.content))
            companies = [row[2] for row in wb["Leads"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(set(companies), {"Alpha", "Gamma"})

            # Filters work when no selection
            with patch.object(main, "StreamingResponse", CapturedResponse):
                res_filtered = main.export_leads(ids=[], state="Maharashtra", db=db)
            wb_filt = load_workbook(BytesIO(res_filtered.content))
            companies_filt = [row[2] for row in wb_filt["Leads"].iter_rows(min_row=2, values_only=True)]
            self.assertEqual(set(companies_filt), {"Alpha", "Beta"})

    def test_incomplete_lead_validation_rules(self) -> None:
        with self.session_factory() as db:
            initial_count = db.scalar(select(func.count()).select_from(Lead))

            # 1. Missing company name raises 422
            with self.assertRaises(HTTPException) as ctx:
                main.create_lead(LeadCreate(company_name="", contact_numbers=["9876543210"]), db)
            self.assertEqual(ctx.exception.status_code, 422)
            self.assertEqual(ctx.exception.detail, "Company name and at least one contact number are required.")

            # 2. Whitespace-only company name raises 422
            with self.assertRaises(HTTPException) as ctx:
                main.create_lead(LeadCreate(company_name="   \t\n  ", contact_numbers=["9876543210"]), db)
            self.assertEqual(ctx.exception.status_code, 422)
            self.assertEqual(ctx.exception.detail, "Company name and at least one contact number are required.")

            # 3. Missing contact number array raises 422
            with self.assertRaises(HTTPException) as ctx:
                main.create_lead(LeadCreate(company_name="Valid Company", contact_numbers=[]), db)
            self.assertEqual(ctx.exception.status_code, 422)
            self.assertEqual(ctx.exception.detail, "Company name and at least one contact number are required.")

            # 4. Blank contact-number array raises 422
            with self.assertRaises(HTTPException) as ctx:
                main.create_lead(LeadCreate(company_name="Valid Company", contact_numbers=["", "   ", " \t "]), db)
            self.assertEqual(ctx.exception.status_code, 422)
            self.assertEqual(ctx.exception.detail, "Company name and at least one contact number are required.")

            # Ensure no partial lead was created on failed validations
            self.assertEqual(db.scalar(select(func.count()).select_from(Lead)), initial_count)

            # 5. Duplicate contact numbers are deduplicated before check and saved
            lead_dedup = main.create_lead(LeadCreate(
                company_name="Dedup Enterprise",
                contact_numbers=["9876543210", "+91 98765 43210", "9876543210"],
            ), db)
            self.assertIsNotNone(lead_dedup.id)
            self.assertEqual(len(lead_dedup.contact_numbers), 1)
            self.assertEqual(lead_dedup.contact_numbers[0], "+91 9876543210")

            # 6. Valid company plus one contact number saves successfully with idempotency
            request = Request({"type": "http", "method": "POST", "path": "/api/leads", "headers": [(b"x-idempotency-key", b"test-key-incomplete-123")]})
            first_user = User(full_name="Valid User", email="user1@example.com", normalized_email="user1@example.com", password_hash="x")
            db.add(first_user)
            db.commit()
            db.refresh(first_user)
            saved1 = main.create_lead(
                LeadCreate(company_name="Valid Tech", contact_numbers=["9811122233"]),
                request, db, first_user
            )
            saved_replay = main.create_lead(
                LeadCreate(company_name="Valid Tech", contact_numbers=["9811122233"]),
                request, db, first_user
            )
            self.assertEqual(saved1.id, saved_replay.id)

            # 7. Existing leads already stored without contact numbers remain readable and exportable
            legacy_lead = Lead(id="legacy-lead-1", company_name="Legacy Preexisting Company")
            db.add(legacy_lead)
            db.commit()

            listed = main.list_leads(page=1, page_size=50, db=db)["leads"]
            found_legacy = next((item for item in listed if item.id == "legacy-lead-1"), None)
            self.assertIsNotNone(found_legacy)
            self.assertEqual(found_legacy.company_name, "Legacy Preexisting Company")

            class CapturedResponse:
                def __init__(self, content, headers=None, **_kwargs):
                    self.content = content.getvalue()
                    self.headers = headers or {}

            with patch.object(main, "StreamingResponse", CapturedResponse):
                res_legacy = main.export_leads(ids=["legacy-lead-1"], db=db)
            wb_legacy = load_workbook(BytesIO(res_legacy.content))
            sheet_legacy = wb_legacy["Leads"]
            self.assertEqual(sheet_legacy["C2"].value, "Legacy Preexisting Company")

if __name__ == "__main__":
    unittest.main()
