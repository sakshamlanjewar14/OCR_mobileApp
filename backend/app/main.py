import csv
import hashlib
import io
import json
import os
import re
import secrets
import logging
import textwrap
import threading
import tempfile
import time
import uuid
import zipfile
import ipaddress
import socket
import urllib.error
import urllib.request
import urllib.parse
from html.parser import HTMLParser
from html import unescape
from datetime import datetime, timedelta, timezone
from pathlib import Path, PurePosixPath
from typing import BinaryIO

from fastapi import BackgroundTasks, Depends, FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from starlette.background import BackgroundTask
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from sqlalchemy import func, inspect, or_, select, text, update, delete
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from PIL import Image, UnidentifiedImageError
try:
    import bcrypt
except ImportError:  # Local minimal environments can still run the auth flow.
    bcrypt = None

try:
    import resource
except ImportError:  # resource is unavailable on Windows development hosts.
    resource = None

from .database import Base, SessionLocal, engine, get_db
from .models import AuthSession, Document, Lead, LeadContact, LeadIdempotency, LeadSourceImage, PasswordResetToken, User
from .ocr import OCRUnavailableError, _parse_fields, extract_fields
from . import paddle_ocr
from .qr import parse_qr_payload
from .geocoding import reverse_geocode
from .schemas import (DocumentRead, DocumentStatusRead, DocumentStatusRequest, ForgotPasswordRequest, LeadCreate,
                      LeadFields, LeadRead, LeadUpdate, LoginRequest, QRParseRequest, RegisterRequest,
                      QRWebsitePreviewRequest, ResetPasswordRequest, ReverseGeocodeRequest, UserRead)
from .translation import warm_translation_model


UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", "./uploads")).resolve()
MAX_UPLOAD_BYTES = max(1, int(os.getenv("MAX_UPLOAD_BYTES", str(10 * 1024 * 1024))))
MAX_BATCH_BYTES = max(MAX_UPLOAD_BYTES, int(os.getenv("MAX_BATCH_BYTES", str(30 * 1024 * 1024))))
MAX_IMAGE_PIXELS = max(1, int(os.getenv("MAX_IMAGE_PIXELS", str(40_000_000))))
MAX_PDF_PAGES = max(1, int(os.getenv("MAX_PDF_PAGES", "5")))
MAX_OCR_CONCURRENCY = max(1, int(os.getenv("MAX_OCR_CONCURRENCY", "1")))
EXPORT_STREAM_THRESHOLD = max(1, int(os.getenv("EXPORT_STREAM_THRESHOLD", "50000")))
SESSION_COOKIE = os.getenv("SESSION_COOKIE", "lead_gennie_session")
SESSION_TTL_HOURS = max(1, int(os.getenv("SESSION_TTL_HOURS", "24")))
RESET_TTL_MINUTES = max(5, int(os.getenv("RESET_TTL_MINUTES", "30")))
IDEMPOTENCY_TTL_HOURS = max(1, int(os.getenv("IDEMPOTENCY_TTL_HOURS", "48")))
AUTH_REQUIRED = os.getenv("AUTH_REQUIRED", "true").lower() not in {"0", "false", "no"}
_PASSWORD_ITERATIONS = 310_000
_ocr_slots = threading.BoundedSemaphore(MAX_OCR_CONCURRENCY)
_active_ocr_jobs = 0
_active_ocr_lock = threading.Lock()
logger = logging.getLogger("lead_gennie.request")
ALLOWED_TYPES = {
    "image/jpeg", "image/png", "image/webp", "image/bmp", "image/tiff", "image/gif", "application/pdf",
}
TYPE_SUFFIXES = {
    "image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp", "image/bmp": ".bmp",
    "image/tiff": ".tiff", "image/gif": ".gif", "application/pdf": ".pdf",
}
DOCUMENT_SUFFIXES = set(TYPE_SUFFIXES) | {".zip"}

app = FastAPI(title="Lead Gennie API", version="0.1.0")
DEFAULT_CORS_ORIGINS = ",".join(
    f"http://{host}:{port}"
    for host in ("localhost", "127.0.0.1")
    for port in range(5173, 5177)
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in os.getenv("CORS_ORIGINS", DEFAULT_CORS_ORIGINS).split(",") if origin.strip()],
    # Authentication uses an HttpOnly session cookie. Browsers reject the
    # credentialed registration/login request unless this is enabled.
    allow_credentials=True,
    allow_methods=["GET", "POST", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["*"],
    expose_headers=["X-Request-ID", "Server-Timing"],
)


@app.middleware("http")
async def measure_request(request, call_next):
    # Accept only UUID-shaped caller IDs; otherwise generate one. Logs contain
    # no request body, card text, filenames, or contact details.
    supplied_id = request.headers.get("x-request-id", "")
    try:
        request_id = str(uuid.UUID(supplied_id))
    except (ValueError, AttributeError):
        request_id = str(uuid.uuid4())
    request.state.request_id = request_id
    started = time.perf_counter()
    status = 500
    try:
        response = await call_next(request)
        status = response.status_code
        response.headers["X-Request-ID"] = request_id
        response.headers["Server-Timing"] = f"app;dur={(time.perf_counter() - started) * 1000:.1f}"
        return response
    finally:
        logger.info("request_complete request_id=%s method=%s path=%s status=%s duration_ms=%.1f",
                    request_id, request.method, request.url.path, status,
                    (time.perf_counter() - started) * 1000)


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _email_key(value: str) -> str:
    return value.strip().casefold()


def _hash_password(password: str) -> str:
    if bcrypt is not None:
        return "bcrypt$" + bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, _PASSWORD_ITERATIONS)
    return f"pbkdf2_sha256${_PASSWORD_ITERATIONS}${salt.hex()}${digest.hex()}"


def _verify_password(password: str, encoded: str) -> bool:
    try:
        if encoded.startswith("bcrypt$") and bcrypt is not None:
            return bcrypt.checkpw(password.encode(), encoded[7:].encode())
        scheme, iterations, salt_hex, digest_hex = encoded.split("$", 3)
        if scheme != "pbkdf2_sha256":
            return False
        digest = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt_hex), int(iterations))
        return secrets.compare_digest(digest.hex(), digest_hex)
    except (ValueError, TypeError):
        return False


def _token_hash(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def _user_read(user: User) -> UserRead:
    return UserRead(id=user.id, full_name=user.full_name, email=user.email, is_verified=user.is_verified)


def _session_response(response, user: User, db: Session):
    raw_token = secrets.token_urlsafe(48)
    db.add(AuthSession(id=secrets.token_hex(32), user_id=user.id, token_hash=_token_hash(raw_token),
                       expires_at=_now() + timedelta(hours=SESSION_TTL_HOURS)))
    db.commit()
    secure = os.getenv("COOKIE_SECURE", "false").lower() == "true"
    same_site = os.getenv("COOKIE_SAMESITE", "none" if secure else "lax").lower()
    if same_site not in {"lax", "strict", "none"}:
        same_site = "none" if secure else "lax"
    # Cross-domain production PWAs require SameSite=None and Secure so the
    # authenticated API requests retain the session cookie. Local HTTP keeps
    # the existing Lax behavior unless explicitly configured otherwise.
    response.set_cookie(SESSION_COOKIE, raw_token, max_age=SESSION_TTL_HOURS * 3600,
                        httponly=True, secure=secure, samesite=same_site, path="/")
    return response


def _current_user(request: Request, db: Session) -> User | None:
    token = request.cookies.get(SESSION_COOKIE)
    if not token:
        return None
    session = db.scalar(select(AuthSession).where(AuthSession.token_hash == _token_hash(token),
                                                  AuthSession.revoked_at.is_(None),
                                                  AuthSession.expires_at > _now()))
    if not session:
        return None
    return db.get(User, session.user_id)


def require_user(request: Request, db: Session = Depends(get_db)) -> User:
    user = _current_user(request, db)
    if not user or not user.is_active:
        raise HTTPException(status_code=401, detail="Authentication required.")
    return user


def _owner_filter(statement, user: User | None, model):
    if isinstance(user, User) and AUTH_REQUIRED:
        return statement.where(model.owner_id == user.id)
    return statement


@app.post("/api/auth/register", response_model=UserRead, status_code=201)
def register(payload: RegisterRequest, db: Session = Depends(get_db)):
    email = _email_key(payload.email)
    if db.scalar(select(User).where(User.normalized_email == email)):
        raise HTTPException(status_code=409, detail="An account with that email already exists.")
    user = User(full_name=payload.full_name.strip(), email=payload.email.strip(), normalized_email=email,
                password_hash=_hash_password(payload.password))
    db.add(user)
    try:
        db.commit()
        db.refresh(user)
    except Exception:
        db.rollback()
        raise HTTPException(status_code=409, detail="An account with that email already exists.")
    response = JSONResponse(_user_read(user).model_dump(), status_code=201)
    return _session_response(response, user, db)


@app.post("/api/auth/login", response_model=UserRead)
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.normalized_email == _email_key(payload.email)))
    now = _now()
    if user and user.locked_until and user.locked_until > now:
        raise HTTPException(status_code=429, detail="Too many attempts. Try again later.", headers={"Retry-After": "900"})
    if not user or not _verify_password(payload.password, user.password_hash):
        if user:
            user.failed_login_attempts += 1
            if user.failed_login_attempts >= 5:
                user.locked_until = now + timedelta(minutes=15)
            db.commit()
        raise HTTPException(status_code=401, detail="Email or password is incorrect.")
    user.failed_login_attempts = 0
    user.locked_until = None
    user.last_login_at = now
    db.commit()
    response = JSONResponse(_user_read(user).model_dump())
    return _session_response(response, user, db)


@app.post("/api/auth/logout")
def logout(request: Request, db: Session = Depends(get_db)):
    token = request.cookies.get(SESSION_COOKIE)
    if token:
        session = db.scalar(select(AuthSession).where(AuthSession.token_hash == _token_hash(token)))
        if session:
            session.revoked_at = _now()
            db.commit()
    response = JSONResponse({"status": "ok"})
    response.delete_cookie(SESSION_COOKIE, path="/")
    return response


@app.get("/api/auth/me", response_model=UserRead)
def me(user: User = Depends(require_user)):
    return _user_read(user)


@app.post("/api/auth/forgot-password")
def forgot_password(payload: ForgotPasswordRequest, db: Session = Depends(get_db)):
    # Deliberately identical response for known and unknown addresses.
    user = db.scalar(select(User).where(User.normalized_email == _email_key(payload.email)))
    response = {"message": "If an account exists, reset instructions have been sent."}
    if user:
        raw_token = secrets.token_urlsafe(40)
        db.add(PasswordResetToken(user_id=user.id, token_hash=_token_hash(raw_token),
                                  expires_at=_now() + timedelta(minutes=RESET_TTL_MINUTES)))
        db.commit()
        if os.getenv("APP_ENV", "development").lower() != "production":
            response["development_reset_token"] = raw_token
    return response


@app.post("/api/auth/reset-password")
def reset_password(payload: ResetPasswordRequest, db: Session = Depends(get_db)):
    reset = db.scalar(select(PasswordResetToken).where(PasswordResetToken.token_hash == _token_hash(payload.token),
                                                       PasswordResetToken.used_at.is_(None),
                                                       PasswordResetToken.expires_at > _now()))
    if not reset:
        raise HTTPException(status_code=400, detail="This reset link is invalid or expired.")
    user = db.get(User, reset.user_id)
    if not user:
        raise HTTPException(status_code=400, detail="This reset link is invalid or expired.")
    user.password_hash = _hash_password(payload.password)
    user.password_changed_at = _now()
    user.failed_login_attempts = 0
    user.locked_until = None
    reset.used_at = _now()
    db.query(AuthSession).filter(AuthSession.user_id == user.id, AuthSession.revoked_at.is_(None)).update({AuthSession.revoked_at: _now()})
    db.commit()
    return {"message": "Password reset successfully. Please log in again."}


@app.on_event("startup")
def initialize_database() -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    Base.metadata.create_all(bind=engine)
    # Add nullable ownership columns to existing databases without touching
    # current OCR/lead data. New records are owned by the authenticated user.
    for table_name in ("leads", "documents"):
        columns = {column["name"] for column in inspect(engine).get_columns(table_name)}
        if "owner_id" not in columns:
            with engine.begin() as connection:
                connection.execute(text(f"ALTER TABLE {table_name} ADD COLUMN owner_id VARCHAR(36)"))
    # ``create_all`` only creates indexes on new tables.  Ensure indexes added
    # for large lead lists are also installed for existing deployments.
    for table in (Lead.__table__, LeadContact.__table__):
        for index in table.indexes:
            index.create(bind=engine, checkfirst=True)
    if engine.dialect.name == "sqlite":
        columns = {column["name"] for column in inspect(engine).get_columns("leads")}
        if "business_description" not in columns:
            with engine.begin() as connection:
                connection.execute(text("ALTER TABLE leads ADD COLUMN business_description TEXT"))
    if engine.dialect.name == "postgresql":
        columns = {column["name"] for column in inspect(engine).get_columns("leads")}
        document_unique_constraints = [
            item["name"] for item in inspect(engine).get_unique_constraints("leads")
            if item.get("column_names") == ["document_id"] and item.get("name")
        ]
        additions = {
            "designation": "VARCHAR(120)",
            "business_description": "TEXT",
            "raw_text": "TEXT",
            "verification_status": "VARCHAR(20) NOT NULL DEFAULT 'verified'",
            "source_card_index": "INTEGER",
            "facebook_link": "VARCHAR(500)",
            "instagram_link": "VARCHAR(500)",
            "whatsapp_link": "VARCHAR(500)",
            "original_ocr_text": "TEXT",
        }
        with engine.begin() as connection:
            for constraint_name in document_unique_constraints:
                quoted_name = constraint_name.replace('"', '""')
                connection.execute(text(f'ALTER TABLE leads DROP CONSTRAINT "{quoted_name}"'))
            for name, definition in additions.items():
                if name not in columns:
                    connection.execute(text(f"ALTER TABLE leads ADD COLUMN {name} {definition}"))
            # Move old JSON arrays and scalar primary contacts into one relational
            # row per value before dropping all legacy contact columns.
            legacy_contact_columns = [name for name in ("contact_numbers", "email_addresses") if name in columns]
            selected_columns = ["id"]
            selected_columns.extend(name for name in ("contact_number", "email") if name in columns)
            selected_columns.extend(name for name in legacy_contact_columns if name not in selected_columns)
            rows = connection.execute(text(
                "SELECT " + ", ".join(selected_columns) + " FROM leads"
            )).mappings().all()
            for row in rows:
                for field, kind, fallback in (
                    ("contact_numbers", "phone", row.get("contact_number")),
                    ("email_addresses", "email", row.get("email")),
                ):
                    raw_values = row.get(field) if field in row else None
                    if isinstance(raw_values, str):
                        try:
                            raw_values = json.loads(raw_values)
                        except json.JSONDecodeError:
                            raw_values = [raw_values]
                    values = raw_values if isinstance(raw_values, list) else []
                    if fallback and fallback not in values:
                        values = [fallback, *values]
                    for position, value in enumerate(dict.fromkeys(str(value).strip() for value in values if value and str(value).strip())):
                        connection.execute(text(
                            "INSERT INTO lead_contacts (id, lead_id, contact_type, value, position) "
                            "VALUES (:id, :lead_id, :contact_type, :value, :position) "
                            "ON CONFLICT (lead_id, contact_type, value) DO NOTHING"
                        ), {"id": str(uuid.uuid4()), "lead_id": row["id"], "contact_type": kind,
                            "value": value, "position": position})
            for name in [column for column in ("contact_numbers", "email_addresses", "contact_number", "email") if column in columns]:
                connection.execute(text(f"ALTER TABLE leads DROP COLUMN {name}"))
            legacy_locations = [name for name in ("city", "district", "state") if name in columns]
            if legacy_locations:
                address_parts = ["NULLIF(BTRIM(company_address), '')"]
                for name in legacy_locations:
                    address_parts.append(
                        f"CASE WHEN NULLIF(BTRIM({name}), '') IS NOT NULL "
                        f"AND POSITION(LOWER(BTRIM({name})) IN LOWER(COALESCE(company_address, ''))) = 0 "
                        f"THEN BTRIM({name}) END"
                    )
                connection.execute(text(
                    "UPDATE leads SET company_address = NULLIF(CONCAT_WS(', ', "
                    + ", ".join(address_parts)
                    + "), '')"
                ))
                for name in legacy_locations:
                    connection.execute(text(f"ALTER TABLE leads DROP COLUMN {name}"))

    # Background tasks live in this local API process. If it was restarted in
    # the middle of OCR, keep the uploaded file and make the interrupted job
    # visible and retryable instead of leaving it stuck as "processing".
    with SessionLocal() as db:
        interrupted = db.scalars(select(Document).where(
            Document.processing_status.in_(["uploaded", "processing"])
        )).all()
        for document in interrupted:
            document.processing_status = "failed"
            document.error_message = "Scanning was interrupted. Retry the scan; the uploaded source file is still available."
        if interrupted:
            db.commit()
    translation_dir = os.getenv("TRANSLATION_MODEL_DIR", "").strip()
    if translation_dir and Path(translation_dir).is_dir():
        threading.Thread(target=warm_translation_model, name="ocr-translation-warmup", daemon=True).start()
    if paddle_ocr.enabled():
        threading.Thread(target=paddle_ocr.warm, name="paddle-ocr-warmup", daemon=True).start()



def _read_document(document: Document) -> DocumentRead:
    linked_leads = {lead.id: lead for lead in document.leads}
    linked_leads.update({source.lead.id: source.lead for source in document.source_links})
    linked_lead = next(iter(linked_leads.values()), None)
    return DocumentRead(
        id=document.id,
        original_filename=document.original_filename,
        content_type=document.content_type,
        processing_status=document.processing_status,
        extracted_fields=document.extracted_fields or {},
        error_message=document.error_message,
        created_at=document.created_at,
        lead_id=linked_lead.id if linked_lead else None,
        lead_count=len(linked_leads),
        resolved_card_indexes=sorted({lead.source_card_index for lead in document.leads if lead.source_card_index is not None}),
    )


def _process_document(document_id: str) -> None:
    db = SessionLocal()
    try:
        document = db.get(Document, document_id)
        if not document:
            return
        document.processing_status = "processing"
        document.error_message = None
        db.commit()
        try:
            created_at = document.created_at
            processing_started_at = datetime.now(timezone.utc)
            if created_at and created_at.tzinfo is None:
                created_at = created_at.replace(tzinfo=timezone.utc)
            queue_ms = max(0, round((processing_started_at - created_at).total_seconds() * 1000)) if created_at else 0
            fields = extract_fields(Path(document.storage_path), document.content_type)
            timing_targets = fields.get("cards") if isinstance(fields.get("cards"), list) else [fields]
            for timing_target in timing_targets:
                if not isinstance(timing_target, dict):
                    continue
                processing_times = timing_target.get("processing_times_ms")
                if not isinstance(processing_times, dict):
                    processing_times = {}
                processing_times["queue_ms"] = queue_ms
                timing_target["processing_times_ms"] = processing_times
            document.extracted_fields = fields
            # OCR results are only candidates. A user must review and save them
            # explicitly so an OCR mismatch cannot silently become a lead.
            document.processing_status = "needs_review"
            candidate = fields.get("cards")
            first_card = candidate[0] if isinstance(candidate, list) and candidate else fields
            has_company = isinstance(first_card, dict) and bool(str(first_card.get("company_name") or "").strip())
            has_other_data = isinstance(first_card, dict) and any(
                first_card.get(key) for key in ("owner_name", "designation", "company_address",
                                                 "contact_numbers", "email_addresses", "website",
                                                 "facebook_link", "instagram_link", "whatsapp_link")
            )
            if has_company:
                document.error_message = None
            elif has_other_data:
                document.error_message = (
                    "No company name was found. This card may contain only QR contact, email, phone, "
                    "or address data. Review the extracted fields and enter the company name manually before saving; "
                    "you can enter the details manually."
                )
            else:
                document.error_message = (
                    "No usable company or contact data was found. Make sure the card text and QR code are clear, "
                    "then retry the scan or enter the details manually."
                )
            db.commit()
        except Exception as exc:  # Preserve the upload and allow manual entry or retry.
            db.rollback()
            document = db.get(Document, document_id)
            if not document:
                return
            document.processing_status = "failed"
            if isinstance(exc, OCRUnavailableError):
                document.error_message = str(exc)
            elif exc.__class__.__name__ == "TesseractNotFoundError":
                document.error_message = "OCR engine is unavailable. Install Tesseract or enter the lead details manually."
            elif isinstance(exc, ValueError):
                document.error_message = str(exc)
            else:
                document.error_message = "Could not read or save this document. Check the file and try again, or enter details manually."
            db.commit()
    finally:
        db.close()


def _content_type_for_header(header: bytes) -> str | None:
    if header.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if header.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if header.startswith(b"%PDF-"):
        return "application/pdf"
    if len(header) >= 12 and header[:4] == b"RIFF" and header[8:12] == b"WEBP":
        return "image/webp"
    if header.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if header.startswith(b"BM"):
        return "image/bmp"
    if header.startswith((b"II*\x00", b"MM\x00*", b"II+\x00", b"MM\x00+")):
        return "image/tiff"
    return None


def _safe_filename(filename: str) -> str:
    return PurePosixPath(filename.replace("\\", "/")).name or "upload"


def _create_document_from_stream(db: Session, filename: str, stream: BinaryIO, owner_id: str | None = None) -> Document:
    filename = _safe_filename(filename)
    header = stream.read(16)
    content_type = _content_type_for_header(header)
    if content_type not in ALLOWED_TYPES:
        raise HTTPException(status_code=415, detail=f"{filename} is not a supported image or PDF file.")
    document_id = str(uuid.uuid4())
    path = UPLOAD_DIR / f"{document_id}{TYPE_SUFFIXES[content_type]}"
    try:
        with path.open("wb") as output:
            output.write(header)
            total = len(header)
            while True:
                chunk = stream.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > MAX_UPLOAD_BYTES:
                    raise HTTPException(status_code=413, detail="File exceeds the configured upload-size limit.")
                output.write(chunk)
        if content_type == "application/pdf":
            import pymupdf
            with pymupdf.open(path) as pdf:
                if pdf.is_encrypted or len(pdf) < 1:
                    raise ValueError("The PDF is encrypted or contains no pages.")
                if len(pdf) > MAX_PDF_PAGES:
                    raise HTTPException(status_code=413, detail=f"PDFs may contain at most {MAX_PDF_PAGES} pages.")
        else:
            try:
                with Image.open(path) as image:
                    if image.width * image.height > MAX_IMAGE_PIXELS:
                        raise HTTPException(status_code=413, detail="Image dimensions exceed the configured pixel limit.")
                    if image.format != {
                        "image/jpeg": "JPEG", "image/png": "PNG", "image/webp": "WEBP",
                        "image/bmp": "BMP", "image/tiff": "TIFF", "image/gif": "GIF",
                    }[content_type]:
                        raise ValueError("The file contents do not match the detected image type.")
                    image.verify()
            except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
                raise ValueError("The image is corrupted or unsupported.") from exc
    except HTTPException:
        path.unlink(missing_ok=True)
        raise
    except Exception as exc:
        path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="The uploaded file is corrupted or invalid.") from exc
    document = Document(
        id=document_id,
        original_filename=filename[:255],
        content_type=content_type,
        storage_path=str(path),
        processing_status="uploaded",
        extracted_fields={},
        owner_id=owner_id,
    )
    db.add(document)
    return document


def _documents_from_zip(db: Session, filename: str, stream: BinaryIO, owner_id: str | None = None) -> list[Document]:
    documents: list[Document] = []
    stream.seek(0)
    try:
        with zipfile.ZipFile(stream) as archive:
            entries = [entry for entry in archive.infolist() if not entry.is_dir()]
            if len(entries) > 25 or sum(entry.file_size for entry in entries) > MAX_BATCH_BYTES:
                raise HTTPException(status_code=413, detail="ZIP exceeds the configured document-count or expanded-size limit.")
            with db.begin_nested():
                for entry in archive.infolist():
                    member_name = entry.filename.replace("\\", "/")
                    safe_name = _safe_filename(member_name)
                    if entry.is_dir() or safe_name.startswith(".") or Path(safe_name).suffix.lower() not in TYPE_SUFFIXES.values():
                        continue
                    with archive.open(entry) as member:
                        try:
                            documents.append(_create_document_from_stream(db, safe_name, member, owner_id))
                        except HTTPException as exc:
                            if exc.status_code != 415:
                                raise
    except Exception as exc:
        for document in documents:
            Path(document.storage_path).unlink(missing_ok=True)
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(status_code=400, detail=f"{filename} could not be read as a valid ZIP archive.") from exc
    if not documents:
        raise HTTPException(status_code=400, detail="No supported images or PDFs were found in the ZIP archive.")
    return documents


def _reserve_ocr_slots(count: int) -> bool:
    reserved = 0
    for _ in range(count):
        if not _ocr_slots.acquire(blocking=False):
            for _ in range(reserved):
                _ocr_slots.release()
            return False
        reserved += 1
    return True


def _process_document_with_request_id(document_id: str, request_id: str, started: float) -> None:
    _process_document(document_id)
    with SessionLocal() as db:
        document = db.get(Document, document_id)
        if document and isinstance(document.extracted_fields, dict):
            fields = dict(document.extracted_fields)
            fields["request_id"] = request_id
            targets = fields.get("cards") if isinstance(fields.get("cards"), list) else [fields]
            for target in targets:
                if isinstance(target, dict):
                    target.setdefault("processing_times_ms", {})["request_id"] = request_id
                    timings = target.get("processing_times_ms", {})
                    logger.info("ocr_complete request_id=%s duration_ms=%.1f queue_ms=%s decode_ms=%s preprocess_ms=%s ocr_ms=%s translation_ms=%s field_extraction_ms=%s qr_ms=%s",
                                request_id, (time.perf_counter() - started) * 1000,
                                timings.get("queue_ms"), timings.get("decode_ms"), timings.get("preprocess_ms"),
                                timings.get("ocr_ms"), timings.get("translation_ms"),
                                timings.get("field_extraction_ms"), timings.get("qr_ms"))
            document.extracted_fields = fields
            db.commit()


def _finish_reserved_job(job, request_id: str) -> None:
    global _active_ocr_jobs
    started = time.perf_counter()
    with _active_ocr_lock:
        _active_ocr_jobs += 1
    try:
        job(started)
    finally:
        with _active_ocr_lock:
            _active_ocr_jobs -= 1
            active_jobs = _active_ocr_jobs
        if resource is not None:
            usage = resource.getrusage(resource.RUSAGE_SELF)
            cpu_seconds = usage.ru_utime + usage.ru_stime
            peak_rss_mb = usage.ru_maxrss / 1024
        else:
            cpu_seconds = peak_rss_mb = 0
        logger.info("ocr_job_complete request_id=%s duration_ms=%.1f active=%s capacity=%s process_cpu_seconds_total=%.2f peak_rss_mb=%s",
                    request_id, (time.perf_counter() - started) * 1000, active_jobs, MAX_OCR_CONCURRENCY, cpu_seconds,
                    f"{peak_rss_mb:.1f}" if peak_rss_mb else "unavailable")
        _ocr_slots.release()


def _process_reserved_document(document_id: str, request_id: str) -> None:
    _finish_reserved_job(lambda started: _process_document_with_request_id(document_id, request_id, started), request_id)


def _process_reserved_documents(document_ids: list[str], request_id: str) -> None:
    def process_batch(started: float) -> None:
        for document_id in document_ids:
            _process_document_with_request_id(document_id, request_id, time.perf_counter())
    _finish_reserved_job(process_batch, request_id)


def _normalize_india_phone(value: str) -> str:
    original = value.strip()
    digits = re.sub(r"\D", "", original)
    if original.startswith("+"):
        return original
    if len(digits) == 10:
        return f"+91 {digits}"
    if len(digits) == 11 and digits.startswith("0") and digits[1] in "6789":
        return f"+91 {digits[1:]}"
    if len(digits) == 12 and digits.startswith("91") and digits[2] in "6789":
        return f"+{digits[:2]} {digits[2:]}"
    if digits.startswith("0091") and len(digits) == 14 and digits[4] in "6789":
        return f"+91 {digits[4:]}"
    return original


def _lead_read(lead: Lead) -> LeadRead:
    result = LeadRead.model_validate(lead)
    result.contact_numbers = [_normalize_india_phone(value) for value in result.contact_numbers]
    result.contact_number = result.contact_numbers[0] if result.contact_numbers else None
    return result


def _filter_lead_statement(
    statement,
    *,
    company_name: str | None = None,
    owner_name: str | None = None,
    city: str | None = None,
    state: str | None = None,
    search: str | None = None,
):
    for column, value in (
        (Lead.company_name, company_name),
        (Lead.owner_name, owner_name),
        (Lead.company_address, city),
        (Lead.company_address, state),
    ):
        if value and value.strip():
            statement = statement.where(column.ilike(f"%{value.strip()}%"))
    if search and search.strip():
        term = f"%{search.strip()}%"
        statement = statement.where(or_(
            Lead.lead_number.ilike(term), Lead.company_name.ilike(term), Lead.owner_name.ilike(term),
            Lead.designation.ilike(term),
            Lead.id.in_(select(LeadContact.lead_id).where(LeadContact.value.ilike(term))),
            Lead.website.ilike(term), Lead.facebook_link.ilike(term), Lead.instagram_link.ilike(term),
            Lead.whatsapp_link.ilike(term), Lead.company_address.ilike(term), Lead.raw_text.ilike(term),
        ))
    return statement


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/ready")
def readiness() -> dict[str, str]:
    """Readiness for routing scans; /health remains a cheap process liveness check."""
    database_ready = False
    try:
        with engine.connect() as connection:
            connection.execute(text("SELECT 1"))
        database_ready = True
    except Exception:
        logger.warning("readiness_failed component=database")
    try:
        from .ocr import _tesseract_available, _uses_windows_ocr_fallback
        ocr_ready = paddle_ocr.available() or _tesseract_available() or _uses_windows_ocr_fallback()
    except Exception:
        ocr_ready = False
    if not ocr_ready:
        logger.warning("readiness_failed component=ocr")
    if not database_ready or not ocr_ready:
        raise HTTPException(status_code=503, detail="Backend dependencies are not ready.")
    return {"status": "ready"}


@app.post("/api/qr/parse")
def parse_qr(payload: QRParseRequest, user: User = Depends(require_user)):
    """Classify scanned QR data locally; this endpoint never opens its URLs."""
    return parse_qr_payload(payload.payload)


class _WebsitePreviewParser(HTMLParser):
    """Small, bounded parser for public card-site metadata; never executes HTML."""
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.description = ""
        self.links: list[str] = []
        self.json_ld: list[str] = []
        self._title_depth = 0
        self._json_depth = 0
        self._buffer: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = {key.lower(): value or "" for key, value in attrs}
        if tag.lower() == "title":
            self._title_depth = 1
            self._buffer = []
        elif self._title_depth:
            self._title_depth += 1
        if tag.lower() == "script" and attributes.get("type", "").lower() == "application/ld+json":
            self._json_depth = 1
            self._buffer = []
        elif self._json_depth:
            self._json_depth += 1
        if tag.lower() == "meta":
            name = (attributes.get("name") or attributes.get("property") or "").lower()
            if name in {"description", "og:description"} and not self.description:
                self.description = attributes.get("content", "").strip()
        if tag.lower() == "a" and attributes.get("href"):
            self.links.append(attributes["href"].strip())

    def handle_data(self, data: str) -> None:
        if self._title_depth or self._json_depth:
            self._buffer.append(data)

    def handle_endtag(self, tag: str) -> None:
        lowered = tag.lower()
        if lowered == "title" and self._title_depth:
            self.title = unescape(" ".join(self._buffer)).strip()[:500]
            self._title_depth = 0
            self._buffer = []
        elif self._title_depth:
            self._title_depth = max(0, self._title_depth - 1)
        if lowered == "script" and self._json_depth:
            self.json_ld.append("".join(self._buffer)[:10000])
            self._json_depth = 0
            self._buffer = []
        elif self._json_depth:
            self._json_depth = max(0, self._json_depth - 1)


def _public_url(url: str) -> str:
    candidate = url.strip()
    if not re.match(r"^[a-z][a-z0-9+.-]*://", candidate, re.I):
        candidate = f"https://{candidate}"
    parsed = urllib.parse.urlparse(candidate)
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise HTTPException(status_code=400, detail="Only public http or https URLs are supported.")
    host = parsed.hostname
    try:
        addresses = {item[4][0] for item in socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)}
    except OSError as exc:
        raise HTTPException(status_code=400, detail="The website host could not be resolved.") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address.split("%", 1)[0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise HTTPException(status_code=400, detail="Private and local network URLs are not allowed.")
    return candidate


@app.post("/api/qr/website-preview")
def qr_website_preview(payload: QRWebsitePreviewRequest, user: User = Depends(require_user)):
    """Read limited public metadata for a QR URL; HTML is never executed."""
    url = _public_url(payload.url)
    class _NoRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, req, fp, code, msg, headers, newurl):
            return None
    opener = urllib.request.build_opener(_NoRedirect())
    request = urllib.request.Request(url, headers={"User-Agent": "LeadGennieBot/1.0", "Accept": "text/html,application/xhtml+xml"})
    try:
        with opener.open(request, timeout=5) as response:
            if response.status >= 300:
                return {"url": url, "fields": {}, "source_text": "", "warnings": ["The website requires a redirect or could not be read."]}
            content_type = response.headers.get_content_type()
            if content_type not in {"text/html", "application/xhtml+xml"}:
                return {"url": url, "fields": {}, "source_text": "", "warnings": ["The URL does not provide a public HTML page."]}
            body = response.read(512_001)
            if len(body) > 512_000:
                return {"url": url, "fields": {}, "source_text": "", "warnings": ["The website response is too large to preview safely."]}
    except (urllib.error.URLError, TimeoutError, ValueError) as exc:
        raise HTTPException(status_code=502, detail="The public website could not be read. Keep the URL for manual review.") from exc
    parser = _WebsitePreviewParser()
    try:
        parser.feed(body.decode("utf-8", errors="replace"))
    except Exception:
        raise HTTPException(status_code=502, detail="The website content could not be parsed safely.")
    fields: dict[str, object] = {"website": url}
    # Many modern sites render their content in JavaScript and expose only a
    # useful title in the initial HTML. Use a conservative title pattern for
    # the company name; never treat a person or address as a company guess.
    if parser.title:
        title_match = re.search(r"\bwith\s+(.+?)\s*:\s*", parser.title, re.IGNORECASE)
        if title_match:
            fields["company_name"] = title_match.group(1).strip(" -|:")[:255]
        elif "|" in parser.title:
            fields["company_name"] = parser.title.split("|", 1)[0].strip()[:255]
        if not parser.description:
            fields["business_description"] = parser.title[:500]
    for link in parser.links:
        absolute = urllib.parse.urljoin(url, link)
        parsed = parse_qr_payload(absolute)
        for key in ("facebook_link", "instagram_link", "whatsapp_link"):
            if parsed.get(key): fields.setdefault(key, parsed[key])
    if parser.description: fields["business_description"] = parser.description
    source_text = " ".join(part for part in (parser.title, parser.description) if part).strip()
    if not fields.get("facebook_link") and not fields.get("instagram_link") and not fields.get("whatsapp_link") and not fields.get("company_name") and not parser.description:
        return {"url": url, "fields": fields, "source_text": source_text, "warnings": ["Only the website URL could be verified. Review other fields manually."]}
    return {"url": url, "fields": fields, "source_text": source_text, "warnings": []}


@app.post("/api/geocoding/reverse")
def resolve_coordinates(payload: ReverseGeocodeRequest, user: User = Depends(require_user)):
    """Explicit caller-triggered lookup; only coordinates go to configured Nominatim."""
    try:
        return reverse_geocode(payload.latitude, payload.longitude)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:
        # Do not return upstream response content, URLs, or request metadata.
        raise HTTPException(status_code=502, detail="Address lookup is temporarily unavailable. Keep or edit the address manually.") from exc


@app.post("/api/documents", response_model=DocumentRead, status_code=202)
async def upload_document(
    request: Request,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_user),
):
    filename = _safe_filename(file.filename or "upload")
    document = _create_document_from_stream(db, filename, file.file, user.id if isinstance(user, User) else None)
    if not _reserve_ocr_slots(1):
        Path(document.storage_path).unlink(missing_ok=True)
        raise HTTPException(status_code=503, detail="OCR is busy. Wait briefly and retry the scan.", headers={"Retry-After": "3"})
    try:
        db.commit()
        db.refresh(document)
    except Exception:
        _ocr_slots.release()
        Path(document.storage_path).unlink(missing_ok=True)
        db.rollback()
        raise
    background_tasks.add_task(_process_reserved_document, document.id, request.state.request_id)
    return _read_document(document)


@app.post("/api/documents/batch", status_code=202)
async def upload_document_batch(
    request: Request,
    background_tasks: BackgroundTasks,
    files: list[UploadFile] = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_user),
):
    documents: list[Document] = []
    skipped_files: list[dict[str, str]] = []
    batch_input_bytes = 0
    for upload in files:
        filename = _safe_filename(upload.filename or "upload")
        try:
            upload.file.seek(0, os.SEEK_END)
            upload_size = upload.file.tell()
            upload.file.seek(0)
            upload_limit = MAX_BATCH_BYTES if Path(filename).suffix.lower() == ".zip" else MAX_UPLOAD_BYTES
            if upload_size > upload_limit:
                raise HTTPException(status_code=413, detail="File exceeds the configured upload-size limit.")
            batch_input_bytes += upload_size
            if batch_input_bytes > MAX_BATCH_BYTES:
                raise HTTPException(status_code=413, detail="Batch exceeds the configured total upload-size limit.")
            if Path(filename).suffix.lower() == ".zip":
                documents.extend(_documents_from_zip(db, filename, upload.file, user.id if isinstance(user, User) else None))
            else:
                documents.append(_create_document_from_stream(db, filename, upload.file, user.id if isinstance(user, User) else None))
        except HTTPException as exc:
            if exc.status_code == 413 and "Batch exceeds" in str(exc.detail):
                for document in documents:
                    Path(document.storage_path).unlink(missing_ok=True)
                db.rollback()
                raise
            skipped_files.append({"filename": filename, "reason": str(exc.detail)})
    if not documents:
        raise HTTPException(status_code=415, detail="No supported image or PDF documents were found.")
    if len(documents) > 25:
        for document in documents:
            Path(document.storage_path).unlink(missing_ok=True)
        db.rollback()
        raise HTTPException(status_code=413, detail="A batch may contain at most 25 documents.")
    if not _reserve_ocr_slots(1):
        for document in documents:
            Path(document.storage_path).unlink(missing_ok=True)
        db.rollback()
        raise HTTPException(status_code=503, detail="OCR is busy. Retry this batch shortly.", headers={"Retry-After": "3"})
    try:
        db.commit()
        for document in documents:
            db.refresh(document)
    except Exception:
        _ocr_slots.release()
        db.rollback()
        for document in documents:
            Path(document.storage_path).unlink(missing_ok=True)
        raise
    background_tasks.add_task(_process_reserved_documents, [document.id for document in documents], request.state.request_id)
    return {"documents": [_read_document(document) for document in documents], "skipped_files": skipped_files}


RAW_FIELD_HEADERS = {
    "company_name": {"company", "company name", "company full name", "full company name", "legal company name", "registered company name", "business", "business name", "firm", "firm name", "organization", "organization name", "organisation", "organisation name"},
    "owner_name": {"owner", "owner name", "owner full name", "contact person", "contact person name", "primary contact name", "contact name", "person name", "full name", "proprietor", "proprietor name", "director name"},
    "designation": {"designation", "job title", "title", "position", "role", "work title"},
    "business_description": {"description", "business description", "company description", "services", "products and services"},
    "company_address": {"address", "company address", "office address", "business address", "registered address", "registered office", "office location", "location"},
    "website": {"website", "web site", "company website", "website url", "web address", "company url", "url", "domain"},
    "facebook_link": {"facebook", "facebook link", "facebook url", "fb link"},
    "instagram_link": {"instagram", "instagram link", "instagram url", "ig link"},
    "whatsapp_link": {"whatsapp link", "whatsapp url", "whatsapp chat link"},
}
RAW_TEXT_HEADERS = {
    "other document text", "other text from document", "other text", "document text", "source text", "raw text",
    "company work data", "company description", "business description", "company profile", "business overview",
}
RAW_PHONE_HEADERS = {
    "phone", "phones", "phone no", "phone nos", "phone number", "phone numbers", "primary phone", "secondary phone", "business phone", "office phone",
    "mobile", "mobiles", "mobile no", "mobile nos", "mobile number", "mobile numbers", "primary mobile", "secondary mobile", "telephone",
    "telephone no", "telephone number", "contact no", "contact nos", "contact number", "contact numbers", "contact phone", "contact mobile",
    "cell", "cell no", "cell number", "whatsapp", "whatsapp no", "whatsapp number",
}
RAW_EMAIL_HEADERS = {
    "email", "emails", "e mail", "email id", "email ids", "email address", "email addresses", "primary email", "secondary email", "business email",
    "work email", "official email", "personal email", "contact email", "office email", "mail id", "mail address",
}


def _normalized_header(header: str) -> str:
    key = re.sub(r"[^a-z0-9]+", " ", header.casefold()).strip()
    return re.sub(r"\s+\d+$", "", key).strip()


def _raw_cell_text(value: object) -> str:
    """Avoid importing an integer phone cell from Excel with a trailing .0."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _raw_row_fields(headers: list[str], values: list[object]) -> dict[str, object]:
    pairs = [(str(header or "").strip(), values[index] if index < len(values) else None)
             for index, header in enumerate(headers)]
    mapped: dict[str, object] = {"contact_numbers": [], "email_addresses": []}
    recognized_pairs: list[tuple[str, object]] = []
    has_company_header = any(_normalized_header(header) in RAW_FIELD_HEADERS["company_name"] for header in headers)
    for header, value in pairs:
        if value in (None, ""):
            continue
        key = _normalized_header(header)
        field = next((field for field, aliases in RAW_FIELD_HEADERS.items() if key in aliases), None)
        if field:
            mapped[field] = _raw_cell_text(value)
            recognized_pairs.append((header, _raw_cell_text(value)))
        elif key in RAW_PHONE_HEADERS:
            mapped["contact_numbers"].extend(part.strip() for part in re.split(r"[\r\n;,/]+", _raw_cell_text(value)) if part.strip())
            recognized_pairs.append((header, _raw_cell_text(value)))
        elif key in RAW_EMAIL_HEADERS:
            mapped["email_addresses"].extend(part.strip() for part in re.split(r"[\r\n;,]+", _raw_cell_text(value)) if part.strip())
            recognized_pairs.append((header, _raw_cell_text(value)))
        elif key in RAW_TEXT_HEADERS:
            recognized_pairs.append((header, _raw_cell_text(value)))
    raw_text = "\n".join(f"{header}: {value}" for header, value in recognized_pairs)
    if not has_company_header:
        parsed = _parse_fields(raw_text)
        for field in ("company_name", "owner_name", "designation", "business_description", "company_address", "website"):
            if parsed.get(field):
                mapped[field] = parsed[field]
        mapped["contact_numbers"].extend(parsed.get("contact_numbers", []))
        mapped["email_addresses"].extend(parsed.get("email_addresses", []))
    mapped["contact_numbers"] = list(dict.fromkeys(mapped["contact_numbers"]))
    mapped["email_addresses"] = list(dict.fromkeys(mapped["email_addresses"]))
    mapped["raw_text"] = raw_text or None
    return mapped


def _find_raw_header_row(rows: list[list[object]], scan_limit: int = 25) -> int | None:
    """Find recognized tabular headers even when a workbook has blank/title rows above them."""
    aliases = [*RAW_FIELD_HEADERS.values(), RAW_TEXT_HEADERS, RAW_PHONE_HEADERS, RAW_EMAIL_HEADERS]
    best_index: int | None = None
    best_score = 0
    for index, row in enumerate(rows[:scan_limit]):
        headers = [str(value or "").strip() for value in row]
        score = sum(
            1 for header in headers
            if any(_normalized_header(header) in field_aliases for field_aliases in aliases)
        )
        if score > best_score:
            best_index, best_score = index, score
    return best_index


def _raw_candidates_from_rows(rows: list[list[object]]) -> list[dict[str, object]]:
    header_index = _find_raw_header_row(rows)
    has_known_headers = header_index is not None
    headers = [str(value or "").strip() for value in rows[header_index]] if has_known_headers else []
    data_start = header_index + 1 if header_index is not None else 0
    candidates = []
    for row_number, row in enumerate(rows[data_start:], start=data_start + 1):
        if not any(value not in (None, "") for value in row):
            continue
        row_headers = headers if has_known_headers else [f"Column {index + 1}" for index in range(len(row))]
        candidates.append({"row": row_number, "fields": _raw_row_fields(row_headers, list(row))})
    return candidates


@app.post("/api/imports/leads/preview")
async def preview_raw_leads(file: UploadFile = File(...), user: User = Depends(require_user)):
    filename = _safe_filename(file.filename or "raw-data")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".txt":
            source = (await file.read()).decode("utf-8-sig")
            fields = _parse_fields(source)
            fields["original_ocr_text"] = source
            fields["raw_text"] = source
            return {"filename": filename, "candidates": [{"row": 1, "fields": fields}]}
        if suffix == ".csv":
            source = (await file.read()).decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(source)))
        elif suffix == ".xlsx":
            workbook = load_workbook(file.file, read_only=True, data_only=True)
            try:
                rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
            finally:
                workbook.close()
        else:
            raise HTTPException(status_code=415, detail="Choose a UTF-8 TXT, CSV, or XLSX raw-data file.")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=400, detail="Text and CSV files must use UTF-8 encoding.") from exc
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not read this raw-data file.") from exc

    if not rows:
        raise HTTPException(status_code=400, detail="The raw-data file is empty.")
    candidates = _raw_candidates_from_rows(rows)
    if not candidates:
        raise HTTPException(status_code=400, detail="No non-empty lead rows were found.")
    return {"filename": filename, "candidates": candidates}


@app.get("/api/documents/pending", response_model=list[DocumentRead])
def list_pending_documents(db: Session = Depends(get_db), user: User = Depends(require_user)):
    statement = select(Document).where(
        Document.processing_status.in_(["uploaded", "processing", "needs_review", "failed"])
    ).order_by(Document.created_at.desc())
    return db.scalars(_owner_filter(statement, user, Document)).all()


@app.post("/api/documents/status", response_model=list[DocumentStatusRead])
def get_document_statuses(payload: DocumentStatusRequest, db: Session = Depends(get_db), user: User = Depends(require_user)):
    # Return compact OCR status during polling without loading extracted text.
    document_ids = list(dict.fromkeys(payload.ids))
    statuses: list[dict[str, object]] = []
    for start in range(0, len(document_ids), 1000):
        chunk = document_ids[start:start + 1000]
        documents = db.scalars(_owner_filter(select(Document).where(Document.id.in_(chunk)), user, Document)).all()
        statuses.extend({
            "id": document.id,
            "processing_status": document.processing_status,
            "error_message": document.error_message,
            "lead_count": len({lead.id for lead in document.leads} | {source.lead.id for source in document.source_links}),
        } for document in documents)
    return statuses


@app.get("/api/documents/{document_id}", response_model=DocumentRead)
def get_document(document_id: str, db: Session = Depends(get_db), user: User = Depends(require_user)):
    document = db.scalar(_owner_filter(select(Document).where(Document.id == document_id), user, Document))
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")
    return _read_document(document)


@app.get("/api/documents/{document_id}/file")
def get_document_file(document_id: str, db: Session = Depends(get_db), user: User = Depends(require_user)):
    document = db.scalar(_owner_filter(select(Document).where(Document.id == document_id), user, Document))
    if not document or not Path(document.storage_path).is_file():
        raise HTTPException(status_code=404, detail="Document file not found.")
    return FileResponse(
        document.storage_path,
        media_type=document.content_type,
        filename=document.original_filename,
        content_disposition_type="inline",
    )


@app.post("/api/documents/{document_id}/retry", response_model=DocumentRead, status_code=202)
def retry_document(document_id: str, request: Request, background_tasks: BackgroundTasks, db: Session = Depends(get_db), user: User = Depends(require_user)):
    document = db.scalar(_owner_filter(select(Document).where(Document.id == document_id), user, Document))
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")
    if document.processing_status != "failed":
        raise HTTPException(status_code=409, detail="Only failed scans can be retried. This scan is already queued, processing, or ready for review.")
    if not Path(document.storage_path).is_file():
        raise HTTPException(status_code=410, detail="The saved source image is unavailable. Upload the card again to retry OCR.")
    if not _reserve_ocr_slots(1):
        raise HTTPException(status_code=503, detail="OCR is busy. Wait briefly and retry the scan.", headers={"Retry-After": "3"})
    try:
        # The conditional transition makes concurrent retry requests idempotent
        # across API workers: only one request can claim a failed document.
        result = db.execute(update(Document).where(
            Document.id == document_id,
            Document.processing_status == "failed",
        ).values(processing_status="uploaded", error_message=None))
        db.commit()
    except Exception:
        _ocr_slots.release()
        db.rollback()
        raise
    if result.rowcount != 1:
        _ocr_slots.release()
        db.rollback()
        raise HTTPException(status_code=409, detail="This scan has already been retried or is no longer eligible for retry.")
    document = db.get(Document, document_id)
    db.refresh(document)
    background_tasks.add_task(_process_reserved_document, document_id, request.state.request_id)
    return _read_document(document)


def _require_english_ocr(values: dict) -> None:
    # `raw_text` is the source/document transcript and may intentionally keep
    # the original Hindi or Marathi text. Only structured lead fields must be
    # Latin-letter values; `original_ocr_text` preserves the untouched source.
    fields = ("company_name", "owner_name", "designation", "business_description", "company_address")
    remaining = [key for key in fields if re.search(r"[\u0900-\u097f]", str(values.get(key) or ""))]
    if remaining:
        raise HTTPException(status_code=422, detail="Enter English/Latin-letter values before saving: " + ", ".join(remaining) + ". Re-scan and translate, or correct the preview manually. Original OCR text is kept separately.")


@app.post("/api/leads", response_model=LeadRead, status_code=201)
def create_lead(payload: LeadCreate, request: Request = None, db: Session = Depends(get_db), user: User = Depends(require_user)):
    # Existing unit callers invoke this function as (payload, db, user). Keep
    # that compatible while FastAPI supplies Request in the HTTP route.
    if request is not None and not isinstance(request, Request):
        legacy_db, legacy_user = request, db
        db, user, request = legacy_db, legacy_user, None
    idempotency_key = request.headers.get("x-idempotency-key", "").strip() if request else ""
    if len(idempotency_key) > 128:
        raise HTTPException(status_code=400, detail="Idempotency key is too long.")
    idempotency_identity = (str(user.id), idempotency_key) if idempotency_key else None
    if idempotency_identity:
        now = _now()
        # Keep the table small while allowing a bounded replay window.
        db.execute(delete(LeadIdempotency).where(LeadIdempotency.expires_at <= now))
        existing_key = db.scalar(select(LeadIdempotency).where(
            LeadIdempotency.user_id == str(user.id),
            LeadIdempotency.idempotency_key == idempotency_key,
            LeadIdempotency.expires_at > now,
        ))
        if existing_key:
            existing = db.get(Lead, existing_key.lead_id)
            if existing:
                return existing
    if (payload.raw_text or payload.document_id) and not payload.source_review_confirmed:
        raise HTTPException(status_code=422, detail="Compare the lead details with the available source before saving.")
    document = None
    if payload.document_id:
        document = db.get(Document, payload.document_id)
        if not document:
            raise HTTPException(status_code=404, detail="Source document not found.")
        if isinstance(user, User) and AUTH_REQUIRED and document.owner_id != user.id:
            raise HTTPException(status_code=404, detail="Source document not found.")
    values = payload.model_dump(exclude={"document_id", "source_document_ids", "source_sides", "document_card_index", "source_review_confirmed"})
    if not values["contact_numbers"] and values["contact_number"]:
        values["contact_numbers"] = [values["contact_number"]]
    if not values["email_addresses"] and values["email"]:
        values["email_addresses"] = [values["email"]]
    values["contact_numbers"] = [_normalize_india_phone(value) for value in values["contact_numbers"]]
    values["contact_number"] = (values["contact_numbers"] or [None])[0]
    values["email"] = (values["email_addresses"] or [None])[0]
    if document and not values["raw_text"]:
        values["raw_text"] = (document.extracted_fields or {}).get("raw_text")
    if document or values.get("original_ocr_text") or payload.source_document_ids:
        _require_english_ocr(values)
    if document and not values.get("original_ocr_text"):
        source = document.extracted_fields or {}
        if payload.document_card_index is not None and source.get("cards"):
            if payload.document_card_index >= len(source["cards"]):
                raise HTTPException(status_code=422, detail="Source card index is invalid.")
            source = source["cards"][payload.document_card_index]
        values["original_ocr_text"] = source.get("original_ocr_text")
    lead = Lead(
        **values,
        document=document,
        source_card_index=payload.document_card_index,
        owner_id=user.id if isinstance(user, User) else None,
        verification_status="verified" if payload.source_review_confirmed or not (payload.raw_text or payload.document_id) else "needs_review",
    )
    db.add(lead)
    source_ids = list(dict.fromkeys(payload.source_document_ids))
    if payload.document_id and payload.document_id not in source_ids:
        source_ids.insert(0, payload.document_id)
    source_documents = {}
    if source_ids:
        source_documents = {item.id: item for item in db.scalars(_owner_filter(select(Document).where(Document.id.in_(source_ids)), user, Document)).all()}
        if len(source_documents) != len(source_ids):
            raise HTTPException(status_code=404, detail="One or more source images could not be found.")
    if document:
        db.flush()
        card_values = (document.extracted_fields or {}).get("cards")
        expected_cards = len(card_values) if isinstance(card_values, list) and card_values else 1
        resolved_indexes = db.scalars(select(Lead.source_card_index).where(
            Lead.document_id == document.id, Lead.source_card_index.is_not(None)
        )).all()
        document.processing_status = (
            "saved" if payload.document_card_index is None or len(set(resolved_indexes)) >= expected_cards else "needs_review"
        )
    if source_ids:
        db.flush()
        for position, source_id in enumerate(source_ids):
            source = source_documents[source_id]
            side = payload.source_sides[position] if position < len(payload.source_sides) and payload.source_sides[position] in {"front", "back"} else ("front" if position == 0 else "back")
            lead.source_images.append(LeadSourceImage(document_id=source_id, side=side, position=position))
            source.processing_status = "saved"
    if idempotency_identity:
        db.flush()
        db.add(LeadIdempotency(
            user_id=str(user.id),
            idempotency_key=idempotency_key[:128],
            lead_id=lead.id,
            expires_at=_now() + timedelta(hours=IDEMPOTENCY_TTL_HOURS),
        ))
    try:
        db.commit()
    except IntegrityError:
        # A second process may have committed the same user/key between the
        # lookup and insert. Its unique constraint is the durable lock; discard
        # this transaction and return the already-created lead.
        db.rollback()
        if idempotency_identity:
            existing_key = db.scalar(select(LeadIdempotency).where(
                LeadIdempotency.user_id == str(user.id),
                LeadIdempotency.idempotency_key == idempotency_key,
                LeadIdempotency.expires_at > _now(),
            ))
            if existing_key:
                existing = db.get(Lead, existing_key.lead_id)
                if existing:
                    return existing
        raise
    db.refresh(lead)
    return lead


@app.get("/api/leads")
def list_leads(
    search: str | None = None,
    company_name: str | None = None,
    owner_name: str | None = None,
    city: str | None = None,
    state: str | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=100),
    db: Session = Depends(get_db),
    user: User = Depends(require_user),
):
    statement = select(Lead).order_by(Lead.created_at.desc(), Lead.id.desc())
    statement = _filter_lead_statement(
        statement, search=search, company_name=company_name, owner_name=owner_name, city=city, state=state,
    )
    statement = _owner_filter(statement, user, Lead)
    total = db.scalar(select(func.count()).select_from(statement.order_by(None).subquery())) or 0
    leads = db.scalars(
        statement.offset((page - 1) * page_size).limit(page_size)
    ).all()
    # Serialize through LeadRead so computed contact arrays (stored in the
    # lead_contacts table) are included just like they are on the detail API.
    return {
        "leads": [_lead_read(lead) for lead in leads],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@app.get("/api/leads/{lead_id}", response_model=LeadRead)
def get_lead(lead_id: str, db: Session = Depends(get_db), user: User = Depends(require_user)):
    lead = db.scalar(_owner_filter(select(Lead).where(Lead.id == lead_id), user, Lead))
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")
    return _lead_read(lead)


@app.patch("/api/leads/{lead_id}", response_model=LeadRead)
def update_lead(lead_id: str, payload: LeadUpdate, db: Session = Depends(get_db), user: User = Depends(require_user)):
    lead = db.scalar(_owner_filter(select(Lead).where(Lead.id == lead_id), user, Lead))
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")
    updates = payload.model_dump(exclude_unset=True)
    source_review_confirmed = updates.pop("source_review_confirmed", False)
    if (updates.get("raw_text", lead.raw_text) or lead.document_id) and not source_review_confirmed:
        raise HTTPException(status_code=422, detail="Compare the lead details with the available source before saving.")
    if "company_name" in updates and updates["company_name"] is None:
        raise HTTPException(status_code=422, detail="Company name is required.")
    if lead.document_id or updates.get("original_ocr_text", lead.original_ocr_text):
        _require_english_ocr({key: updates.get(key, getattr(lead, key)) for key in ("company_name", "owner_name", "designation", "business_description", "company_address", "raw_text")})
    contact_fields = {"contact_number", "contact_numbers", "email", "email_addresses"}
    for key, value in updates.items():
        if key not in contact_fields:
            setattr(lead, key, value)
    if "contact_numbers" in updates:
        lead.contact_numbers = [_normalize_india_phone(value) for value in (updates["contact_numbers"] or [])]
    elif "contact_number" in updates:
        lead.contact_numbers = [_normalize_india_phone(updates["contact_number"])] if updates["contact_number"] else []
    if "email_addresses" in updates:
        lead.email_addresses = updates["email_addresses"] or []
    elif "email" in updates:
        lead.email_addresses = [updates["email"]] if updates["email"] else []
    if source_review_confirmed:
        lead.verification_status = "verified"
    db.commit()
    db.refresh(lead)
    return _lead_read(lead)


@app.post("/api/leads/batch", status_code=201)
def create_leads_batch(payload: list[LeadFields], db: Session = Depends(get_db), user: User = Depends(require_user)):
    """Explicit one-click import; leave imported rows flagged for user review."""
    if not payload:
        raise HTTPException(status_code=400, detail="No valid lead rows were provided.")
    # Flush bounded chunks while keeping one transaction.  This allows large
    # spreadsheet imports to use predictable ORM memory without exposing
    # partially committed imports if a later row fails validation.
    batch_size = 1000
    created_count = 0
    try:
        for start in range(0, len(payload), batch_size):
            for item in payload[start:start + batch_size]:
                values = item.model_dump()
                if not values["contact_numbers"] and values["contact_number"]:
                    values["contact_numbers"] = [values["contact_number"]]
                if not values["email_addresses"] and values["email"]:
                    values["email_addresses"] = [values["email"]]
                values["contact_numbers"] = [_normalize_india_phone(value) for value in values["contact_numbers"]]
                values["contact_number"] = (values["contact_numbers"] or [None])[0]
                values["email"] = (values["email_addresses"] or [None])[0]
                db.add(Lead(**values, verification_status="needs_review",
                            owner_id=user.id if isinstance(user, User) else None))
            db.flush()
            created_count += min(batch_size, len(payload) - start)
            db.expunge_all()
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"created": created_count}


@app.delete("/api/leads/{lead_id}", status_code=204)
def delete_lead(lead_id: str, db: Session = Depends(get_db), user: User = Depends(require_user)):
    lead = db.scalar(_owner_filter(select(Lead).where(Lead.id == lead_id), user, Lead))
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")
    if lead.document and len(lead.document.leads) <= 1:
        lead.document.processing_status = "needs_review"
    db.delete(lead)
    db.commit()


EXPORT_LEAD_HEADERS = [
    "S.No.", "Autogenerated ID", "Company name", "Owner name", "Designation", "Company address", "Data check",
]
EXPORT_END_HEADERS = ["Website", "Facebook link", "Instagram link", "WhatsApp link", "Business description", "Other document text"]


@app.get("/api/exports/leads")
def export_leads(
    search: str | None = None,
    ids: list[str] = Query(default=[]),
    company_name: str | None = None,
    owner_name: str | None = None,
    city: str | None = None,
    state: str | None = None,
    export_count: int | None = Query(default=None, ge=1),
    db: Session = Depends(get_db),
    user: User = Depends(require_user),
):
    statement = select(Lead).order_by(Lead.created_at.desc(), Lead.id.desc())
    if ids:
        statement = statement.where(Lead.id.in_(ids))
    else:
        statement = _filter_lead_statement(
            statement, search=search, company_name=company_name, owner_name=owner_name, city=city, state=state,
        )
    statement = _owner_filter(statement, user, Lead)
    selected_id_query = statement.with_only_columns(Lead.id, maintain_column_froms=True).order_by(
        Lead.created_at.desc(), Lead.id.desc()
    )
    if export_count is not None:
        selected_id_query = selected_id_query.limit(export_count)
    selected_ids = selected_id_query.subquery()
    selected_count = db.scalar(select(func.count()).select_from(selected_ids)) or 0
    contact_counts = (
        select(LeadContact.lead_id, LeadContact.contact_type, func.count().label("value_count"))
        .where(LeadContact.lead_id.in_(select(selected_ids.c.id)))
        .group_by(LeadContact.lead_id, LeadContact.contact_type)
        .subquery()
    )
    maximums = db.execute(
        select(contact_counts.c.contact_type, func.max(contact_counts.c.value_count))
        .group_by(contact_counts.c.contact_type)
    ).all()
    maximum_by_type = {contact_type: count for contact_type, count in maximums}
    phone_column_count = max(1, maximum_by_type.get("phone", 0))
    email_column_count = max(1, maximum_by_type.get("email", 0))
    headers = [
        *EXPORT_LEAD_HEADERS,
        *(f"Contact Number {index}" for index in range(1, phone_column_count + 1)),
        *(f"Email ID {index}" for index in range(1, email_column_count + 1)),
        *EXPORT_END_HEADERS,
    ]
    streaming_export = selected_count > EXPORT_STREAM_THRESHOLD
    workbook = Workbook(write_only=streaming_export)
    sheet = workbook.create_sheet("Leads") if streaming_export else workbook.active
    if not streaming_export:
        sheet.title = "Leads"
    sheet.freeze_panes = "A2"
    header_fill = PatternFill(fill_type="solid", fgColor="226B4B")
    header_font = Font(bold=True, color="FFFFFF")
    column_widths = [len(header) for header in headers]
    for column, header in enumerate(headers, start=1):
        cell = WriteOnlyCell(sheet, value=header) if streaming_export else sheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        if streaming_export:
            # Write-only worksheets accept a complete row rather than cell
            # assignment; collect the styled header cells below.
            continue
    if streaming_export:
        header_cells = []
        for header in headers:
            cell = WriteOnlyCell(sheet, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
            header_cells.append(cell)
        sheet.append(header_cells)
    sheet.row_dimensions[1].height = 30
    export_query = statement.with_only_columns(
        Lead.id, Lead.lead_number, Lead.company_name, Lead.owner_name, Lead.designation,
        Lead.company_address, Lead.verification_status, Lead.website, Lead.facebook_link,
        Lead.instagram_link, Lead.whatsapp_link, Lead.business_description, Lead.raw_text, Lead.created_at,
        maintain_column_froms=True,
    ).where(Lead.id.in_(select(selected_ids.c.id))).order_by(Lead.created_at.desc(), Lead.id.desc())
    row_count = 0
    page_size = 500
    last_created_at = None
    last_id = None
    while True:
        page_query = export_query
        if last_created_at is not None and last_id is not None:
            page_query = page_query.where(or_(
                Lead.created_at < last_created_at,
                (Lead.created_at == last_created_at) & (Lead.id < last_id),
            ))
        rows = db.execute(page_query.limit(page_size)).all()
        if not rows:
            break
        lead_ids = [row.id for row in rows]
        contact_rows = db.execute(
            select(LeadContact.lead_id, LeadContact.contact_type, LeadContact.value)
            .where(LeadContact.lead_id.in_(lead_ids))
            .order_by(LeadContact.lead_id, LeadContact.position, LeadContact.id)
        ).all()
        contacts_by_lead: dict[str, tuple[list[str], list[str]]] = {
            lead_id: ([], []) for lead_id in lead_ids
        }
        for lead_id, contact_type, value in contact_rows:
            contacts = contacts_by_lead[lead_id]
            contacts[0 if contact_type == "phone" else 1].append(value)
        for lead_id, (phone_values, email_values) in contacts_by_lead.items():
            contacts_by_lead[lead_id] = ([_normalize_india_phone(value) for value in phone_values], email_values)
        for row_offset, row in enumerate(rows):
            contact_numbers, email_addresses = contacts_by_lead[row.id]
            values = [
                row_count + row_offset + 1, row.lead_number, row.company_name, row.owner_name, row.designation, row.company_address, row.verification_status,
                *contact_numbers, *([None] * (phone_column_count - len(contact_numbers))),
                *email_addresses, *([None] * (email_column_count - len(email_addresses))),
                row.website, row.facebook_link, row.instagram_link, row.whatsapp_link, row.business_description, row.raw_text,
            ]
            safe_values = [values[0], *(str(value) if value is not None else "" for value in values[1:])]
            sheet.append(safe_values)
            for index, raw_value in enumerate(safe_values):
                value = str(raw_value) if raw_value is not None else ""
                if value:
                    longest_line = max((len(line) for line in value.splitlines()), default=0)
                    column_widths[index] = max(column_widths[index], longest_line)
                if not streaming_export:
                    cell = sheet.cell(row=sheet.max_row, column=index + 1)
                    # Explicit string cells keep phone numbers (including +91) exact
                    # while preventing spreadsheet formula execution.
                    if value.lstrip().startswith(("=", "+", "-", "@")):
                        cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        row_count += len(rows)
        last_created_at = rows[-1].created_at
        last_id = rows[-1].id
    for index, longest in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(60, max(14, longest + 2))
    if not streaming_export:
        for row_index in range(2, sheet.max_row + 1):
            visible_lines = 1
            for column_index, width in enumerate(column_widths, start=1):
                value = sheet.cell(row=row_index, column=column_index).value
                if value:
                    usable_width = max(1, min(60, max(14, width + 2)) - 2)
                    visible_lines = max(visible_lines, sum(
                        max(1, len(textwrap.wrap(line, width=usable_width, break_long_words=True, break_on_hyphens=True)))
                        for line in str(value).splitlines()
                    ))
            sheet.row_dimensions[row_index].height = min(390, max(20, visible_lines * 15))
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_count + 1}"
    from io import BytesIO

    if streaming_export:
        output_path = tempfile.NamedTemporaryFile(prefix="lead-export-", suffix=".xlsx", delete=False).name
        workbook.save(output_path)
        output_file = open(output_path, "rb")
        # Stream bounded chunks; iterating a binary file directly asks
        # Starlette to read newline-delimited records, which can stall on a ZIP
        # member containing no newline for a long time.
        output = iter(lambda: output_file.read(64 * 1024), b"")
        background = BackgroundTask(lambda: (output_file.close(), Path(output_path).unlink(missing_ok=True)))
    else:
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        background = None
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="leads.xlsx"'},
        background=background,
    )
